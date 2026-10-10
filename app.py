import os
import re
import csv
import html
import math
import hashlib
import time
import tempfile
import unicodedata
from datetime import datetime
from io import BytesIO
from difflib import SequenceMatcher

import json
import random
# PyMuPDF; "import fitz" is deprecated in recent releases.
import pymupdf as fitz
import numpy as np
import streamlit as st
from dotenv import load_dotenv
from openai import AsyncOpenAI
from google import genai
from google.genai import types
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ragas.llms import llm_factory
import asyncio
from ragas.embeddings.base import BaseRagasEmbedding
from ragas.metrics.collections import (
    Faithfulness,
    AnswerRelevancy,
    ContextPrecision,
    ContextRecall
)

# ---------------------------------------------------
# SETTINGS
# ---------------------------------------------------

load_dotenv()

API_KEY = os.getenv("GEMINI_API_KEY")

GENERATION_MODEL = "gemini-3.5-flash-lite"
EMBEDDING_MODEL = "gemini-embedding-2"
EMBEDDING_DIMENSION = 768

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200
TOP_K = 4


# ---------------------------------------------------
# STREAMLIT
# ---------------------------------------------------

st.set_page_config(
    page_title="LearnLens AI",
    page_icon="🎓",
    layout="wide"
)

# ---------------------------------------------------
# UI DESIGN SYSTEM (presentation only)
# All styling lives here. HTML is used only for static,
# hard-coded text; user or model content is always rendered
# with native Streamlit elements (never injected into HTML).
# ---------------------------------------------------

GLOBAL_CSS = """
<style>
:root {
    --ll-navy: #0B1B3F;
    --ll-indigo: #4F46E5;
    --ll-violet: #7C3AED;
    --ll-cyan: #06B6D4;
    --ll-muted: #64748B;
    --ll-line: rgba(79, 70, 229, 0.14);
    --ll-tint: rgba(79, 70, 229, 0.045);
}

/* Layout */
[data-testid="stMainBlockContainer"] {
    max-width: 1180px;
    padding-top: 1.6rem;
    padding-bottom: 4rem;
}

/* Typography */
h1, h2, h3, h4 { letter-spacing: -0.01em; }
[data-testid="stHeading"] h3 {
    font-size: 1.2rem;
    font-weight: 700;
    padding-top: 0.6rem;
}
[data-testid="stCaptionContainer"] { color: var(--ll-muted); }

/* Hero */
.ll-hero {
    position: relative;
    overflow: hidden;
    border-radius: 18px;
    padding: 28px 32px 22px 32px;
    margin-bottom: 0.9rem;
    background: linear-gradient(135deg, #0B1B3F 0%, #1E1B4B 58%, #312E81 100%);
    color: #F8FAFC;
}
.ll-hero-title {
    font-size: clamp(1.7rem, 3vw, 2.35rem);
    font-weight: 800;
    line-height: 1.15;
    color: #FFFFFF;
}
.ll-hero-tagline {
    margin-top: 0.35rem;
    font-size: 1.05rem;
    font-weight: 600;
    color: #C7D2FE;
}
.ll-hero-copy {
    margin: 0.55rem 0 0.9rem 0;
    max-width: 720px;
    color: #CBD5E1;
    font-size: 0.97rem;
}
.ll-chip {
    display: inline-block;
    margin: 0 6px 6px 0;
    padding: 4px 12px;
    border-radius: 999px;
    font-size: 0.8rem;
    font-weight: 600;
    color: #E0E7FF;
    background: rgba(255, 255, 255, 0.08);
    border: 1px solid rgba(255, 255, 255, 0.18);
}
.ll-chip-cyan {
    color: #CFFAFE;
    border-color: rgba(103, 232, 249, 0.45);
}
.ll-powered {
    position: absolute;
    top: 16px;
    right: 22px;
    font-size: 0.75rem;
    color: #A5B4FC;
}
@media (max-width: 700px) {
    .ll-hero { padding: 22px 20px 16px 20px; }
    .ll-powered { position: static; display: block; margin-bottom: 8px; }
}

/* Section headers */
.ll-section { margin: 0.2rem 0 0.8rem 0; }
.ll-section-title { font-size: 1.45rem; font-weight: 750; line-height: 1.25; }
.ll-section-sub { margin-top: 3px; font-size: 0.95rem; color: var(--ll-muted); }
.ll-eyebrow {
    margin: 0.2rem 0 0.25rem 0;
    font-size: 0.72rem;
    font-weight: 700;
    letter-spacing: 0.09em;
    text-transform: uppercase;
    color: var(--ll-indigo);
}

/* Tabs as compact navigation */
[data-testid="stTabs"] [data-baseweb="tab-list"] {
    gap: 4px;
    padding: 4px;
    border-radius: 12px;
    background: var(--ll-tint);
    border: 1px solid var(--ll-line);
    overflow-x: auto;
    flex-wrap: nowrap;
    scrollbar-width: thin;
}
[data-testid="stTabs"] button[role="tab"] {
    height: auto;
    padding: 6px 12px;
    border-radius: 9px;
    white-space: nowrap;
}
[data-testid="stTabs"] button[role="tab"] p {
    font-size: 0.88rem;
    font-weight: 600;
}
[data-testid="stTabs"] button[role="tab"][aria-selected="true"] {
    background: #FFFFFF;
    box-shadow: 0 1px 4px rgba(15, 23, 42, 0.14);
}
[data-testid="stTabs"] button[role="tab"][aria-selected="true"] p {
    color: var(--ll-indigo);
}
[data-testid="stTabs"] [data-baseweb="tab-highlight"],
[data-testid="stTabs"] [data-baseweb="tab-border"] { display: none; }
[data-testid="stTabs"] [data-baseweb="tab-panel"] { padding-top: 1.1rem; }

/* Metric cards */
[data-testid="stMetric"] {
    padding: 12px 14px;
    border-radius: 12px;
    background: var(--ll-tint);
    border: 1px solid var(--ll-line);
}
[data-testid="stMetricLabel"] p {
    font-size: 0.74rem;
    font-weight: 700;
    letter-spacing: 0.05em;
    text-transform: uppercase;
    color: var(--ll-muted);
    white-space: normal;
}
[data-testid="stMetricValue"] {
    font-size: 1.35rem;
    font-weight: 750;
}
[data-testid="stMetricValue"] > div {
    white-space: normal;
    overflow-wrap: anywhere;
    line-height: 1.25;
}

/* Buttons */
.stButton > button {
    border-radius: 10px;
    font-weight: 600;
    transition: box-shadow 0.15s ease;
}
.stButton > button[kind="primary"],
button[data-testid="stBaseButton-primary"] {
    border: none;
    color: #FFFFFF;
    background: linear-gradient(135deg, #4F46E5 0%, #7C3AED 100%);
}
.stButton > button[kind="primary"]:hover,
button[data-testid="stBaseButton-primary"]:hover {
    box-shadow: 0 4px 14px rgba(79, 70, 229, 0.32);
}

/* Expanders, uploaders, alerts */
[data-testid="stExpander"] details {
    border-radius: 12px;
    border: 1px solid var(--ll-line);
}
[data-testid="stExpander"] summary p { font-weight: 600; }
[data-testid="stFileUploaderDropzone"] {
    border-radius: 12px;
    border: 1.5px dashed rgba(79, 70, 229, 0.38);
    background: var(--ll-tint);
}
[data-testid="stAlertContainer"] { border-radius: 12px; }

/* Cards: any container whose key starts with "card_" */
[class*="st-key-card_"] {
    padding: 1rem 1.1rem;
    border-radius: 16px;
    background: var(--ll-tint);
    border: 1px solid var(--ll-line);
}
.st-key-card_tutor_answer {
    border-left: 4px solid var(--ll-indigo);
}
.st-key-card_tutor_refusal {
    border-left: 4px solid var(--ll-cyan);
    background: rgba(6, 182, 212, 0.07);
    border-color: rgba(6, 182, 212, 0.30);
}
[class*="st-key-card_question"] [role="radiogroup"] { gap: 6px; }
[class*="st-key-card_question"] [role="radiogroup"] label {
    width: 100%;
    padding: 8px 10px;
    border-radius: 10px;
    border: 1px solid var(--ll-line);
}
[class*="st-key-card_misconception"] {
    border-left: 4px solid #F59E0B;
}
.st-key-card_synthetic {
    border: 1px dashed rgba(124, 58, 237, 0.45);
    background: rgba(124, 58, 237, 0.05);
}

/* Sidebar */
[data-testid="stSidebar"] {
    border-right: 1px solid rgba(15, 23, 42, 0.06);
}
[data-testid="stSidebar"] [data-testid="stMetric"] { padding: 8px 10px; }
[data-testid="stSidebar"] [data-testid="stMetricValue"] { font-size: 1.1rem; }
</style>
"""

st.markdown(GLOBAL_CSS, unsafe_allow_html=True)


def html_text(text):
    """Escapes text for the few static HTML snippets."""

    return html.escape(str(text), quote=True)


def section_header(title, subtitle=None, eyebrow=None):
    """Consistent page/section heading (static text only)."""

    parts = ['<div class="ll-section">']

    if eyebrow:
        parts.append(
            f'<div class="ll-eyebrow">{html_text(eyebrow)}</div>'
        )

    parts.append(
        f'<div class="ll-section-title">{html_text(title)}</div>'
    )

    if subtitle:
        parts.append(
            f'<div class="ll-section-sub">{html_text(subtitle)}</div>'
        )

    parts.append("</div>")

    st.markdown("".join(parts), unsafe_allow_html=True)


def eyebrow_label(text):
    """Small uppercase label such as ANSWER / SOURCES USED."""

    st.markdown(
        f'<div class="ll-eyebrow">{html_text(text)}</div>',
        unsafe_allow_html=True
    )


DIFFICULTY_BADGES = {
    "Easy": ("blue", ":material/signal_cellular_alt_1_bar:"),
    "Medium": ("violet", ":material/signal_cellular_alt_2_bar:"),
    "Hard": ("orange", ":material/signal_cellular_alt:")
}


def difficulty_badge(difficulty):

    color, icon = DIFFICULTY_BADGES.get(
        difficulty,
        ("gray", None)
    )

    st.badge(str(difficulty), icon=icon, color=color)


def status_badges(labels):
    """Row of green ✓ badges (e.g. indexing status)."""

    with st.container(horizontal=True, gap="small"):

        for label in labels:
            st.badge(
                str(label),
                icon=":material/check_circle:",
                color="green"
            )


# ---------------------------------------------------
# HERO
# ---------------------------------------------------

st.markdown(
    """
<div class="ll-hero">
  <div class="ll-powered">Powered by Gemini</div>
  <div class="ll-hero-title">🎓 LearnLens AI</div>
  <div class="ll-hero-tagline">Your AI-powered personalized learning companion</div>
  <div class="ll-hero-copy">
    Turn PDFs, lecture slides, videos and images into a source-grounded,
    adaptive learning experience.
  </div>
  <span class="ll-chip">Multimodal Learning</span>
  <span class="ll-chip ll-chip-cyan">Source-Grounded AI</span>
  <span class="ll-chip">Adaptive Personalization</span>
</div>
""",
    unsafe_allow_html=True
)

with st.expander("🚀 3-Minute Demo Flow", expanded=False):

    st.markdown(
        """
1. **🧭 Learner Profile** — set level, goal and explanation style.
2. **📄 PDF · 📽️ Slides · 🎞️ Video** — upload learning materials
   (they join one shared knowledge base).
3. **🧭 Learner Profile → Diagnostic** — 3 questions
   (Easy → Medium → Hard) estimate your starting level.
4. **🎓 Unified Tutor** — ask an in-material question (cited answer),
   then an off-material one (LearnLens declines instead of guessing).
5. **🧠 Adaptive Assessment** — answer questions; difficulty adapts.
6. **📊 Mastery Dashboard** — mastery, weak topics, likely misconceptions.
7. **🧪 Evaluation** — RAGAS metrics, Quick Benchmark, synthetic
   personalization benchmark.
"""
    )


# ---------------------------------------------------
# GEMINI
# ---------------------------------------------------

if not API_KEY:
    st.error(
        "Gemini API key could not be found. Copy `.env.example` "
        "to `.env`, set `GEMINI_API_KEY`, and restart the app."
    )
    st.stop()

client = genai.Client(api_key=API_KEY)

# ---------------------------------------------------
# PERSISTENT LEARNER STATE
# ---------------------------------------------------

LEARNER_STATE_FILE = os.path.join(
    os.path.dirname(
        os.path.abspath(__file__)
    ),
    "learner_state.json"
)


def save_learner_state():

    state_data = {
        "learner_profile": st.session_state.get(
            "learner_profile",
            {
                "level": "Intermediate",
                "goal": "",
                "explanation_style": "Balanced",
                "diagnostic_completed": False,
                "diagnostic_score": 0,
                "diagnostic_total": 0
            }
        ),
        "quiz_score": st.session_state.get(
            "quiz_score",
            0
        ),
        "quiz_total": st.session_state.get(
            "quiz_total",
            0
        ),
        "quiz_difficulty": st.session_state.get(
            "quiz_difficulty",
            "Medium"
        ),
        "quiz_correct_streak": st.session_state.get(
            "quiz_correct_streak",
            0
        ),
        "quiz_question_number": st.session_state.get(
            "quiz_question_number",
            0
        ),
        "topic_performance": st.session_state.get(
            "topic_performance",
            {}
        ),
        "quiz_history": st.session_state.get(
            "quiz_history",
            []
        ),
        "learner_mastery": st.session_state.get(
            "learner_mastery",
            {}
        ),
        "conversation_counts": st.session_state.get(
            "conversation_counts",
            {}
        )
    }

    temporary_file = (
        LEARNER_STATE_FILE
        + ".tmp"
    )

    # A locked file (e.g. OneDrive sync) must never break
    # a quiz submit or tutor answer: warn and continue.
    try:

        with open(
            temporary_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                state_data,
                file,
                ensure_ascii=False,
                indent=4
            )

        os.replace(
            temporary_file,
            LEARNER_STATE_FILE
        )

        return True

    except Exception as e:

        st.warning(
            "Learner progress could not be saved to "
            f"learner_state.json ({type(e).__name__}). "
            "It is kept for this session."
        )

        return False


def load_learner_state():

    if not os.path.exists(
        LEARNER_STATE_FILE
    ):
        return

    try:

        with open(
            LEARNER_STATE_FILE,
            "r",
            encoding="utf-8"
        ) as file:

            state_data = json.load(
                file
            )

        persistent_keys = [
            "learner_profile",
            "quiz_score",
            "quiz_total",
            "quiz_difficulty",
            "quiz_correct_streak",
            "quiz_question_number",
            "topic_performance",
            "quiz_history",
            "learner_mastery",
            "conversation_counts"
        ]

        dict_keys = [
            "learner_profile",
            "topic_performance",
            "learner_mastery",
            "conversation_counts"
        ]

        for key in persistent_keys:

            if key in state_data:

                # Skip values with an unexpected type so a
                # hand-edited state file cannot crash the app.
                if (
                    key in dict_keys
                    and not isinstance(state_data[key], dict)
                ):
                    continue

                if (
                    key == "quiz_history"
                    and not isinstance(state_data[key], list)
                ):
                    continue

                st.session_state[
                    key
                ] = state_data[
                    key
                ]

        # These values belong only to the current screen/run.
        st.session_state.quiz_question = None
        st.session_state.quiz_answered = False
        st.session_state.seen_questions = set()
        st.session_state.last_answer_correct = False
        st.session_state.last_feedback = ""

    except Exception as e:

        st.warning(
            f"Saved learner state could not be loaded: {e}"
        )


if "persistent_state_loaded" not in st.session_state:

    load_learner_state()

    st.session_state.persistent_state_loaded = True


def ensure_core_state():
    """
    Guarantees every core learner/assessment key exists with a
    valid type, whatever learner_state.json contained (missing
    keys, wrong types, old format). Runs on every rerun; it only
    fills or repairs values, it never clears valid history.
    """

    expected_types = {
        "quiz_score": (int, 0),
        "quiz_total": (int, 0),
        "quiz_correct_streak": (int, 0),
        "quiz_question_number": (int, 0),
        "topic_performance": (dict, {}),
        "quiz_history": (list, []),
        "learner_mastery": (dict, {}),
        "conversation_counts": (dict, {})
    }

    for key, (expected_type, default_value) in expected_types.items():

        value = st.session_state.get(key)

        if (
            not isinstance(value, expected_type)
            or isinstance(value, bool)
        ):
            st.session_state[key] = (
                default_value.copy()
                if isinstance(default_value, (dict, list))
                else default_value
            )

    if st.session_state.quiz_score > st.session_state.quiz_total:
        st.session_state.quiz_score = st.session_state.quiz_total

    if st.session_state.get("quiz_difficulty") not in (
        "Easy",
        "Medium",
        "Hard"
    ):
        st.session_state.quiz_difficulty = "Medium"

    # Topic stats must have numeric correct/total (total > 0);
    # malformed entries are dropped instead of crashing views.
    st.session_state.topic_performance = {
        str(topic): {
            "correct": int(stats.get("correct", 0)),
            "total": int(stats.get("total", 0))
        }
        for topic, stats in st.session_state.topic_performance.items()
        if isinstance(stats, dict)
        and isinstance(stats.get("total"), int)
        and stats.get("total", 0) > 0
        and isinstance(stats.get("correct", 0), int)
    }

    st.session_state.learner_mastery = {
        str(topic): min(1.0, max(0.0, float(value)))
        for topic, value in st.session_state.learner_mastery.items()
        if isinstance(value, (int, float))
        and not isinstance(value, bool)
    }

    st.session_state.quiz_history = [
        item for item in st.session_state.quiz_history
        if isinstance(item, dict)
    ]

    st.session_state.setdefault("quiz_question", None)
    st.session_state.setdefault("quiz_answered", False)
    st.session_state.setdefault("seen_questions", set())
    st.session_state.setdefault("last_answer_correct", False)
    st.session_state.setdefault("last_feedback", "")


ensure_core_state()


# ---------------------------------------------------
# PDF RAG FUNCTIONS
# ---------------------------------------------------

def split_text(text, chunk_size=1200, overlap=200):
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        piece = text[start:end].strip()

        if piece:
            chunks.append(piece)

        if end == len(text):
            break

        start = end - overlap

    return chunks


def create_file_signature(files):
    hasher = hashlib.sha256()

    for file_name, file_bytes in files:
        hasher.update(file_name.encode("utf-8"))
        hasher.update(file_bytes)

    return hasher.hexdigest()


PDF_MAX_IMAGES_PER_PAGE = 2
PDF_MIN_IMAGE_SIDE = 80

GEMINI_IMAGE_MIME_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp"
}


def extract_pdf_page_images(document, page):
    """
    Extracts up to PDF_MAX_IMAGES_PER_PAGE embedded figures
    from one PDF page. Very small images (icons, logos,
    decorations) are skipped.
    """

    images = []

    try:
        page_image_list = page.get_images(full=True)
    except Exception:
        return images

    for image_info in page_image_list:

        if len(images) >= PDF_MAX_IMAGES_PER_PAGE:
            break

        xref = image_info[0]

        try:
            extracted = document.extract_image(xref)
        except Exception:
            continue

        if not extracted:
            continue

        width = extracted.get("width", 0)
        height = extracted.get("height", 0)

        if (
            width < PDF_MIN_IMAGE_SIDE
            or height < PDF_MIN_IMAGE_SIDE
        ):
            continue

        extension = str(
            extracted.get("ext", "")
        ).lower()

        mime_type = GEMINI_IMAGE_MIME_TYPES.get(
            extension
        )

        image_bytes = extracted.get("image")

        if mime_type is None:

            # Convert formats Gemini does not accept
            # (e.g. JPX, JBIG2) to PNG.
            try:
                pixmap = fitz.Pixmap(document, xref)

                if pixmap.n - pixmap.alpha >= 4:
                    pixmap = fitz.Pixmap(
                        fitz.csRGB,
                        pixmap
                    )

                image_bytes = pixmap.tobytes("png")
                mime_type = "image/png"

            except Exception:
                continue

        if not image_bytes:
            continue

        images.append({
            "bytes": image_bytes,
            "mime_type": mime_type,
            "width": width,
            "height": height
        })

    return images


def collect_source_image_parts(
    sources,
    source_types=("pdf", "slide"),
    max_per_source=2
):
    """
    Builds Gemini image parts for retrieved sources.
    Each PDF page / slide contributes at most
    max_per_source images, and the same page is never
    sent twice (several chunks can share one page).
    """

    parts = []
    seen_pages = set()

    for source in sources:

        source_type = source.get(
            "source_type",
            "pdf"
        )

        if source_type not in source_types:
            continue

        page_key = (
            source.get("file", ""),
            source_type,
            source.get("page", source.get("slide", ""))
        )

        if page_key in seen_pages:
            continue

        seen_pages.add(page_key)

        page_images = source.get(
            "images",
            []
        )[:max_per_source]

        if not page_images:
            continue

        if source_type == "slide":
            label = (
                f"[{source.get('file', '')}, "
                f"Slide {source.get('slide', source.get('page', ''))}]"
            )
        else:
            label = (
                f"[{source.get('file', '')}, "
                f"Page {source.get('page', '')}]"
            )

        label_added = False

        for image_data in page_images:

            try:

                image_part = types.Part.from_bytes(
                    data=image_data["bytes"],
                    mime_type=image_data["mime_type"]
                )

                if not label_added:
                    parts.append(
                        f"Image(s) from {label}:"
                    )
                    label_added = True

                parts.append(image_part)

            except Exception:
                pass

    return parts


def extract_pdf_chunks(files):
    chunks = []
    full_text_parts = []
    file_infos = []

    for file_name, file_bytes in files:
        document = fitz.open(
            stream=file_bytes,
            filetype="pdf"
        )

        file_infos.append({
            "name": file_name,
            "pages": len(document)
        })

        for page_number, page in enumerate(document):
            page_text = page.get_text("text").strip()

            page_images = extract_pdf_page_images(
                document,
                page
            )

            if not page_text:

                if not page_images:
                    continue

                # Figure-only page: keep it retrievable so
                # its images can still reach Gemini.
                page_text = (
                    f"Page {page_number + 1} contains "
                    "visual content (figure/image) with "
                    "little or no extractable text."
                )

            full_text_parts.append(
                f"\n===== {file_name} / Page {page_number + 1} =====\n"
                + page_text
            )

            page_chunks = split_text(
                page_text,
                CHUNK_SIZE,
                CHUNK_OVERLAP
            )

            for chunk_number, chunk_text in enumerate(page_chunks):
                chunks.append({
                    "text": chunk_text,
                    "file": file_name,
                    "page": page_number + 1,
                    "chunk": chunk_number + 1,
                    "images": page_images
                })

    return chunks, "\n".join(full_text_parts), file_infos


def create_embeddings(chunks):
    vectors = []
    batch_size = 40

    for start in range(0, len(chunks), batch_size):
        batch = chunks[start:start + batch_size]

        contents = [
            types.Content(
                parts=[
                    types.Part.from_text(text=chunk["text"])
                ]
            )
            for chunk in batch
        ]

        result = client.models.embed_content(
            model=EMBEDDING_MODEL,
            contents=contents,
            config=types.EmbedContentConfig(
                output_dimensionality=EMBEDDING_DIMENSION
            )
        )

        for embedding in result.embeddings:
            vectors.append(embedding.values)

    return np.array(vectors, dtype=np.float32)


def create_question_embedding(question):
    result = client.models.embed_content(
        model=EMBEDDING_MODEL,
        contents=question,
        config=types.EmbedContentConfig(
            output_dimensionality=EMBEDDING_DIMENSION
        )
    )

    return np.array(
        result.embeddings[0].values,
        dtype=np.float32
    )


def cosine_similarity(query_vector, document_vectors):
    query_norm = np.linalg.norm(query_vector)
    document_norms = np.linalg.norm(document_vectors, axis=1)

    similarities = (
        document_vectors @ query_vector
    ) / (
        document_norms * query_norm + 1e-10
    )

    return similarities


def retrieve_chunks(question, chunks, document_embeddings, top_k=4):
    question_embedding = create_question_embedding(question)

    similarities = cosine_similarity(
        question_embedding,
        document_embeddings
    )

    top_indices = np.argsort(similarities)[::-1][:top_k]

    results = []

    for index in top_indices:
        result = chunks[index].copy()
        result["score"] = float(similarities[index])
        results.append(result)

    return results


# ---------------------------------------------------
# IMAGE QA FUNCTION
# ---------------------------------------------------

def answer_image_question(uploaded_image, question):
    image_bytes = uploaded_image.getvalue()
    mime_type = uploaded_image.type

    image_part = types.Part.from_bytes(
        data=image_bytes,
        mime_type=mime_type
    )

    prompt = f"""
You are LearnLens AI, a multimodal personalized learning assistant.

The user uploaded an image and asked a question.

Rules:
1. Answer based on the visual content of the uploaded image.
2. If the answer is uncertain, clearly say so.
3. Explain clearly and educationally.
4. Answer in the same language as the user's question.

USER QUESTION:
{question}
"""

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=[image_part, prompt]
    )

    return response.text

# ---------------------------------------------------
# VIDEO FUNCTIONS
# ---------------------------------------------------

def upload_video_to_gemini(uploaded_video):
    """
    Uploads a video to Gemini Files API.
    Reuses the same uploaded video during the session.
    """

    video_bytes = uploaded_video.getvalue()

    signature = hashlib.sha256(
        video_bytes
    ).hexdigest()

    if (
        st.session_state.get("video_signature") == signature
        and "video_file_ref" in st.session_state
    ):
        return st.session_state.video_file_ref

    suffix = os.path.splitext(
        uploaded_video.name
    )[1]

    if not suffix:
        suffix = ".mp4"

    temp_path = None

    try:
        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=suffix
        ) as temp_file:

            temp_file.write(video_bytes)
            temp_path = temp_file.name

        video_file = client.files.upload(
            file=temp_path
        )

        # Wait until Gemini finishes processing
        # (bounded, so a stuck upload cannot hang the app).
        processing_deadline = time.time() + 600

        while True:

            state = getattr(
                video_file.state,
                "name",
                str(video_file.state)
            )

            if state == "ACTIVE":
                break

            if state == "FAILED":
                raise RuntimeError(
                    "Gemini could not process the video."
                )

            if time.time() > processing_deadline:
                raise RuntimeError(
                    "Gemini video processing timed out "
                    "(10 minutes). Try a shorter video."
                )

            time.sleep(2)

            video_file = client.files.get(
                name=video_file.name
            )

        file_ref = {
            "name": video_file.name,
            "uri": video_file.uri,
            "mime_type": video_file.mime_type
        }

        st.session_state.video_signature = signature
        st.session_state.video_file_ref = file_ref

        return file_ref

    finally:

        if (
            temp_path
            and os.path.exists(temp_path)
        ):
            os.remove(temp_path)


def answer_video_question(video_ref, question):

    video_part = types.Part.from_uri(
        file_uri=video_ref["uri"],
        mime_type=video_ref["mime_type"]
    )

    prompt = f"""
You are LearnLens AI, a personalized multimodal
learning assistant.

Analyze the uploaded educational video and answer
the student's question.

RULES:

1. Base your answer on the video.
2. Do not invent information.
3. If possible, include relevant timestamps.
4. Explain the answer clearly like a tutor.
5. Answer in the same language as the student's question.
6. If the requested information is not present in the video,
   clearly say so.

STUDENT QUESTION:

{question}
"""

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=[
            video_part,
            prompt
        ]
    )

    return response.text

def create_video_knowledge_chunks(video_ref, video_name):
    """
    Converts an uploaded video into timestamped
    knowledge-base chunks.
    """

    video_part = types.Part.from_uri(
        file_uri=video_ref["uri"],
        mime_type=video_ref["mime_type"]
    )

    prompt = """
You are preparing an educational video for a
multimodal learning knowledge base.

Analyze the entire video and divide it into meaningful
educational segments.

For every segment provide:
- start timestamp
- end timestamp
- main topic
- concise but informative description of what is taught

IMPORTANT:
- Preserve important definitions, explanations and examples.
- Do not invent information.
- Use timestamps from the video.
- Return ONLY valid JSON.
- Do not use Markdown.

Return exactly this structure:

{
  "segments": [
    {
      "start": "00:00",
      "end": "00:20",
      "topic": "Topic name",
      "text": "Educational content explained in this segment"
    }
  ]
}
"""

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=[
            video_part,
            prompt
        ],
        config=types.GenerateContentConfig(
            response_mime_type="application/json"
        )
    )

    cleaned = clean_json_response(
        response.text
    )

    data = json.loads(cleaned)

    segments = data.get(
        "segments",
        []
    )

    chunks = []

    for segment in segments:

        # Empty segments cannot be embedded or cited.
        if (
            not isinstance(segment, dict)
            or not str(segment.get("text", "")).strip()
        ):
            continue

        chunks.append({
            "text": segment.get(
                "text",
                ""
            ),
            "file": video_name,
            "source_type": "video",
            "start": segment.get(
                "start",
                "Unknown"
            ),
            "end": segment.get(
                "end",
                "Unknown"
            ),
            "topic": segment.get(
                "topic",
                "General"
            )
        })

    return chunks


# ---------------------------------------------------
# POWERPOINT / SLIDE FUNCTIONS
# ---------------------------------------------------

def extract_pptx_slides(file_name, file_bytes):
    """
    Extracts text and embedded images from PowerPoint slides.
    """

    presentation = Presentation(
        BytesIO(file_bytes)
    )

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        texts = []
        images = []

        for shape in slide.shapes:

            # Extract text
            if hasattr(shape, "text"):

                shape_text = shape.text.strip()

                if shape_text:
                    texts.append(shape_text)

                        # Extract table content
            if getattr(shape, "has_table", False):

                table_rows = []

                for row in shape.table.rows:

                    row_values = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    table_rows.append(
                        " | ".join(row_values)
                    )

                table_text = "\n".join(
                    table_rows
                ).strip()

                if table_text:
                    texts.append(
                        table_text
                    )

            # Extract embedded pictures
            if (
                shape.shape_type
                == MSO_SHAPE_TYPE.PICTURE
            ):

                try:

                    image = shape.image

                    images.append({
                        "bytes": image.blob,
                        "mime_type": image.content_type
                    })

                except Exception:
                    pass

        slide_text = "\n".join(texts).strip()

        if not slide_text:

            slide_text = (
                f"Slide {slide_number} contains "
                "visual content with little or no "
                "extractable text."
            )

        slides.append({
            "text": slide_text,
            "file": file_name,
            "page": slide_number,
            "slide": slide_number,
            "images": images
        })

    return slides


def answer_slide_question(
    question,
    slide_chunks,
    slide_embeddings
):

    retrieved_slides = retrieve_chunks(
        question,
        slide_chunks,
        slide_embeddings,
        min(3, len(slide_chunks))
    )

    context_parts = []

    for number, slide in enumerate(
        retrieved_slides,
        start=1
    ):

        context_parts.append(
            f"""
SOURCE {number}

File: {slide['file']}
Slide: {slide['slide']}

{slide['text']}
"""
        )

    context = "\n".join(context_parts)

    prompt = f"""
You are LearnLens AI, a multimodal personalized
learning assistant.

Answer the student's question using ONLY the
retrieved PowerPoint slides.

RULES:

1. Use only the provided slide content and images.
2. Do not invent information.
3. Answer in the same language as the student.
4. Explain clearly like a tutor.
5. Cite relevant slides using:

   [File name, Slide X]

6. If the information cannot be found, say so.

RETRIEVED SLIDES:

{context}

STUDENT QUESTION:

{question}
"""

    contents = [prompt]

    # Add relevant images from retrieved slides.
    # Maximum 2 images per slide to keep requests manageable.
    for slide in retrieved_slides:

        for image_data in slide["images"][:2]:

            try:

                contents.append(
                    types.Part.from_bytes(
                        data=image_data["bytes"],
                        mime_type=image_data[
                            "mime_type"
                        ]
                    )
                )

            except Exception:
                pass

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=contents
    )

    return response.text, retrieved_slides

# ---------------------------------------------------
# UNIFIED MULTIMODAL KNOWLEDGE BASE
# ---------------------------------------------------

def build_unified_knowledge_base():

    items = []
    vectors = []

    # -----------------------------
    # PDF SOURCES
    # -----------------------------

    if (
        "chunks" in st.session_state
        and "embeddings" in st.session_state
    ):

        for chunk, vector in zip(
            st.session_state.chunks,
            st.session_state.embeddings
        ):

            item = chunk.copy()

            item["source_type"] = "pdf"

            item["citation"] = (
                f"[{item['file']}, "
                f"Page {item['page']}]"
            )

            items.append(item)
            vectors.append(vector)

    # -----------------------------
    # SLIDE SOURCES
    # -----------------------------

    if (
        "slide_chunks" in st.session_state
        and "slide_embeddings" in st.session_state
    ):

        for slide, vector in zip(
            st.session_state.slide_chunks,
            st.session_state.slide_embeddings
        ):

            item = slide.copy()

            item["source_type"] = "slide"

            item["citation"] = (
                f"[{item['file']}, "
                f"Slide {item['slide']}]"
            )

            items.append(item)
            vectors.append(vector)

    # -----------------------------
    # VIDEO SOURCES
    # -----------------------------

    if (
        "video_chunks" in st.session_state
        and "video_embeddings" in st.session_state
    ):

        for video_chunk, vector in zip(
            st.session_state.video_chunks,
            st.session_state.video_embeddings
        ):

            item = video_chunk.copy()

            item["source_type"] = "video"

            item["citation"] = (
                f"[{item['file']}, "
                f"{item['start']} - "
                f"{item['end']}]"
            )

            items.append(item)
            vectors.append(vector)

    if not items:
        return [], None

    # Attach cached topic / subtopic / key concept /
    # prerequisite tags (no API call here).
    apply_knowledge_tags(items)

    vectors = np.array(
        vectors,
        dtype=np.float32
    )

    return items, vectors


def retrieve_unified_sources(
    question,
    items,
    vectors,
    top_k=6
):

    question_embedding = (
        create_question_embedding(
            question
        )
    )

    similarities = cosine_similarity(
        question_embedding,
        vectors
    )

    top_indices = np.argsort(
        similarities
    )[::-1][:top_k]

    results = []

    for index in top_indices:

        item = items[index].copy()

        item["score"] = float(
            similarities[index]
        )

        results.append(item)

    return results


def answer_unified_question(
    question,
    retrieved_sources,
    learner_profile=None,
    topic_performance=None
):
    

    if learner_profile is None:
        learner_profile = {}

    if topic_performance is None:
        topic_performance = {}

    learner_level = learner_profile.get(
        "level",
        "Intermediate"
    )

    learning_goal = learner_profile.get(
        "goal",
        ""
    )

    explanation_style = learner_profile.get(
        "explanation_style",
        "Balanced"
    )

    # -------------------------------------------
    # FIND WEAK TOPICS
    # -------------------------------------------

    weak_topics = []

    for topic, stats in topic_performance.items():

        total = stats.get(
            "total",
            0
        )

        correct = stats.get(
            "correct",
            0
        )

        if total > 0:

            accuracy = correct / total

            if accuracy < 0.70:

                weak_topics.append(
                    (
                        topic,
                        accuracy
                    )
                )

    # EMA learner mastery below 0.50 also marks a weak topic
    # (covers diagnostic and tutor-conversation evidence).
    accuracy_weak_names = {
        item[0] for item in weak_topics
    }

    for topic, mastery in st.session_state.get(
        "learner_mastery",
        {}
    ).items():

        try:
            mastery = float(mastery)
        except (TypeError, ValueError):
            continue

        if (
            mastery < 0.50
            and topic not in accuracy_weak_names
        ):
            weak_topics.append((topic, mastery))

    weak_topics.sort(
        key=lambda item: item[1]
    )

    weak_topic_names = [
        item[0]
        for item in weak_topics[:3]
    ]

    if weak_topic_names:

        weak_topics_text = ", ".join(
            weak_topic_names
        )

    else:

        weak_topics_text = (
            "No weak topic has been identified yet."
        )

    # -------------------------------------------
    # LEVEL PERSONALIZATION
    # -------------------------------------------

    level_instructions = {
        "Beginner": (
            "Use simple language. Explain essential terms "
            "before using them. Break difficult ideas into "
            "small steps and avoid unnecessary complexity."
        ),
        "Intermediate": (
            "Use balanced academic language. Explain the "
            "main idea clearly and include relevant reasoning "
            "or examples when supported by the material."
        ),
        "Advanced": (
            "Use more technical and precise language. Focus "
            "on relationships, reasoning, comparisons, and "
            "deeper interpretation when supported by the "
            "learning material."
        )
    }

    style_instructions = {
        "Concise": (
            "Keep the answer compact and focused on the "
            "essential information."
        ),
        "Balanced": (
            "Give enough explanation to understand the "
            "concept without unnecessary detail."
        ),
        "Detailed": (
            "Provide a more detailed step-by-step explanation "
            "while remaining grounded in the source."
        )
    }

    level_instruction = level_instructions.get(
        learner_level,
        level_instructions["Intermediate"]
    )

    style_instruction = style_instructions.get(
        explanation_style,
        style_instructions["Balanced"]
    )

    # -------------------------------------------
    # BUILD SOURCE CONTEXT
    # -------------------------------------------

    context_parts = []

    for number, source in enumerate(
        retrieved_sources,
        start=1
    ):

        source_type = source[
            "source_type"
        ]

        context_parts.append(
            f"""
SOURCE {number}

Type: {source_type}
Citation: {source["citation"]}

Content:
{source["text"]}
"""
        )

    context = "\n".join(
        context_parts
    )

    # -------------------------------------------
    # PERSONALIZED TUTOR PROMPT
    # -------------------------------------------

    prompt = f"""
You are LearnLens AI, a personalized multimodal
learning tutor.

The student has uploaded learning materials such as
PDFs, lecture slides and videos.

A retrieval system selected the most relevant content
from the unified knowledge base.

LEARNER PROFILE:

Current level:
{learner_level}

Learning goal:
{learning_goal if learning_goal else "No specific goal provided"}

Preferred explanation style:
{explanation_style}

Topics currently needing more practice:
{weak_topics_text}

PERSONALIZATION INSTRUCTIONS:

1. Adapt the complexity of the explanation to the
   student's current level.

2. Follow this level instruction:
   {level_instruction}

3. Follow this explanation-style instruction:
   {style_instruction}

4. When relevant to the student's question, connect the
   explanation to the student's learning goal.

5. If the question relates to a weak topic, give slightly
   more explanatory support.

6. Personalization must NEVER cause you to add facts that
   are not supported by the retrieved learning material.

STRICT GROUNDING RULES:

1. Answer ONLY using the retrieved material.
2. Never invent course information.
3. If the material does not contain enough information,
   clearly say:
   "Bu bilgi yüklenen ders materyallerinde bulunamadı."
4. Answer in the same language as the student.
5. Explain concepts clearly like a tutor.
6. Every important claim must contain its source citation.
7. Preserve the exact citation supplied below.
8. PDF sources must cite page numbers.
9. Slides must cite slide numbers.
10. Videos must cite timestamps.
11. If several sources support the answer, combine them.

RETRIEVED MATERIAL:

{context}

STUDENT QUESTION:

{question}
"""

    contents = [
        prompt
    ]

    # Include images from relevant slides and PDF pages
    # (max 2 per slide/page, each page sent once).
    contents.extend(
        collect_source_image_parts(
            retrieved_sources,
            source_types=("slide", "pdf"),
            max_per_source=2
        )
    )

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=contents
    )

    return response.text

class FixedGeminiRagasEmbeddings(
    BaseRagasEmbedding
):

    def embed_text(
        self,
        text,
        **kwargs
    ):

        result = client.models.embed_content(
            model=EMBEDDING_MODEL,
            contents=text,
            config=types.EmbedContentConfig(
                output_dimensionality=EMBEDDING_DIMENSION
            )
        )

        return list(
            result.embeddings[0].values
        )

    async def aembed_text(
        self,
        text,
        **kwargs
    ):

        return await asyncio.to_thread(
            self.embed_text,
            text
        )

def run_ragas_evaluation(
    question,
    answer,
    retrieved_contexts,
    reference_answer
):

    ragas_client = AsyncOpenAI(
        api_key=API_KEY,
        base_url=(
            "https://generativelanguage.googleapis.com/"
            "v1beta/openai/"
        )
    )

    evaluator_llm = llm_factory(
        GENERATION_MODEL,
        provider="openai",
        client=ragas_client
    )

    evaluator_embeddings = (
        FixedGeminiRagasEmbeddings()
    )

    faithfulness_metric = Faithfulness(
        llm=evaluator_llm
    )

    answer_relevancy_metric = AnswerRelevancy(
        llm=evaluator_llm,
        embeddings=evaluator_embeddings
    )

    context_precision_metric = ContextPrecision(
        llm=evaluator_llm
    )

    context_recall_metric = ContextRecall(
        llm=evaluator_llm
    )

    faithfulness_result = (
        faithfulness_metric.score(
            user_input=question,
            response=answer,
            retrieved_contexts=retrieved_contexts
        )
    )

    answer_relevancy_result = (
        answer_relevancy_metric.score(
            user_input=question,
            response=answer
        )
    )

    context_precision_result = (
        context_precision_metric.score(
            user_input=question,
            reference=reference_answer,
            retrieved_contexts=retrieved_contexts
        )
    )

    context_recall_result = (
        context_recall_metric.score(
            user_input=question,
            reference=reference_answer,
            retrieved_contexts=retrieved_contexts
        )
    )

    return {
        "faithfulness": float(
            faithfulness_result.value
        ),
        "answer_relevancy": float(
            answer_relevancy_result.value
        ),
        "context_precision": float(
            context_precision_result.value
        ),
        "context_recall": float(
            context_recall_result.value
        )
    }

# ---------------------------------------------------
# ADAPTIVE QUIZ FUNCTIONS
# ---------------------------------------------------

def clean_json_response(text):
    text = text.strip()

    if text.startswith("```json"):
        text = text[7:]

    if text.startswith("```"):
        text = text[3:]

    if text.endswith("```"):
        text = text[:-3]

    return text.strip()


def generate_quiz_question(chunks, difficulty):
    """
    Generates one source-grounded quiz question
    from one PDF chunk.
    """

    source_chunk = random.choice(chunks)

    difficulty_instructions = {
        "Easy": """
Create a simple recall or basic understanding question.
The answer should be directly visible in the material.
""",
        "Medium": """
Create an understanding or interpretation question.
The student should need to think about the material,
not only copy a sentence.
""",
        "Hard": """
Create a more challenging application, reasoning,
comparison, or inference question based strictly
on the provided material.
"""
    }

    prompt = f"""
You are LearnLens AI.

Create ONE multiple-choice quiz question from the
learning material below.

Difficulty: {difficulty}

{difficulty_instructions[difficulty]}

IMPORTANT RULES:

1. Use ONLY the provided learning material.
2. Create exactly 4 answer choices.
3. Only one answer must be correct.
4. Do not use outside information.
5. Write the question in the same language as the material.
6. Return ONLY valid JSON.
7. Do not use Markdown.
8. correct_index must be 0, 1, 2, or 3.

Return this exact JSON structure:

{{
    "question": "question text",
    "options": [
        "option 1",
        "option 2",
        "option 3",
        "option 4"
    ],
    "correct_index": 0,
    "explanation": "short explanation",
    "topic": "short topic name"
}}

LEARNING MATERIAL:

{source_chunk["text"]}
"""

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=prompt
    )

    cleaned = clean_json_response(response.text)

    quiz_data = json.loads(cleaned)

    if len(quiz_data["options"]) != 4:
        raise ValueError("Quiz must contain exactly four options.")

    correct_index = int(quiz_data["correct_index"])

    if correct_index not in [0, 1, 2, 3]:
        raise ValueError("Invalid correct answer index.")

    quiz_data["correct_index"] = correct_index

    quiz_data["source_file"] = source_chunk["file"]
    quiz_data["source_page"] = source_chunk["page"]
    quiz_data["difficulty"] = difficulty

    return quiz_data


def update_difficulty(is_correct):
    """
    Adjusts question difficulty according to performance.
    """

    if is_correct:

        st.session_state.quiz_correct_streak += 1

        if st.session_state.quiz_correct_streak >= 2:

            if st.session_state.quiz_difficulty == "Easy":
                st.session_state.quiz_difficulty = "Medium"

            elif st.session_state.quiz_difficulty == "Medium":
                st.session_state.quiz_difficulty = "Hard"

            st.session_state.quiz_correct_streak = 0

    else:

        st.session_state.quiz_correct_streak = 0

        if st.session_state.quiz_difficulty == "Hard":
            st.session_state.quiz_difficulty = "Medium"

        elif st.session_state.quiz_difficulty == "Medium":
            st.session_state.quiz_difficulty = "Easy"


def update_topic_performance(topic, is_correct):
    """
    Tracks performance for each topic.
    """

    if "topic_performance" not in st.session_state:
        st.session_state.topic_performance = {}

    topic = canonical_topic_name(topic)

    if topic not in st.session_state.topic_performance:

        st.session_state.topic_performance[topic] = {
            "correct": 0,
            "total": 0
        }

    st.session_state.topic_performance[topic]["total"] += 1

    if is_correct:
        st.session_state.topic_performance[topic]["correct"] += 1


def reset_quiz():
    """
    Resets only the current assessment session.
    Long-term learner history (topic_performance,
    quiz_history, learner_mastery, conversation_counts,
    seen_questions) is kept across assessments.
    """

    st.session_state.quiz_score = 0
    st.session_state.quiz_total = 0

    diagnostic_completed = st.session_state.get(
        "learner_profile",
        {}
    ).get(
        "diagnostic_completed",
        False
    )

    # Keep the diagnostic-based starting difficulty.
    if (
        not diagnostic_completed
        or "quiz_difficulty" not in st.session_state
    ):
        st.session_state.quiz_difficulty = "Medium"

    st.session_state.quiz_correct_streak = 0

    st.session_state.quiz_question = None

    st.session_state.quiz_answered = False

    st.session_state.quiz_question_number = 0

    st.session_state.last_answer_correct = False

    st.session_state.last_feedback = ""

    # First-run initialization only; never clears history.
    st.session_state.setdefault("topic_performance", {})
    st.session_state.setdefault("quiz_history", [])
    st.session_state.setdefault("learner_mastery", {})
    st.session_state.setdefault("conversation_counts", {})
    st.session_state.setdefault("seen_questions", set())


def normalize_question_text(text):

    return " ".join(
        text.lower().strip().split()
    )

def is_similar_question(
    new_question,
    old_question,
    threshold=0.82
):

    new_text = normalize_question_text(
        new_question
    )

    old_text = normalize_question_text(
        old_question
    )

    if new_text == old_text:
        return True

    similarity = SequenceMatcher(
        None,
        new_text,
        old_text
    ).ratio()

    return similarity >= threshold

def get_assessment_pool(source_scope):

    items, vectors = build_unified_knowledge_base()

    if not items:
        return [], None

    filtered_items = []
    filtered_vectors = []

    for index, item in enumerate(items):

        include = False

        if source_scope == "All Materials":
            include = True

        elif (
            source_scope == "PDF"
            and item["source_type"] == "pdf"
        ):
            include = True

        elif (
            source_scope == "Slides"
            and item["source_type"] == "slide"
        ):
            include = True

        elif (
            source_scope == "Video"
            and item["source_type"] == "video"
        ):
            include = True

        if include:

            filtered_items.append(item)

            filtered_vectors.append(
                vectors[index]
            )

    if not filtered_items:
        return [], None

    return (
        filtered_items,
        np.array(
            filtered_vectors,
            dtype=np.float32
        )
    )


def verify_assessment_question(
    question_data,
    context,
    relevant_sources=None
):

    verification_prompt = f"""
You are an assessment verifier.

Check whether the generated educational question is
fully supported by the provided learning material.

Verify ALL of the following:

1. The question can be answered using ONLY the material.
2. The correct answer is actually correct.
3. The explanation matches the material.
4. The question does not require outside knowledge.
5. The source citation is supported by the material.
6. For MCQ, exactly one option must be correct.
7. For Numerical questions, the numerical answer must be
   derivable from the material.
8. The question must not be ambiguous.
9. The question must be completely self-contained.
10. Reject the question if it refers to an example, figure,
    table, slide, image, previous text, or numerical data that
    is not explicitly reproduced inside the question.
11. A student must be able to answer without opening the
    original learning material.
12. The student sees ONLY the question text and options.
    No image, chart, figure, equation image or table is shown.
    Any attached images are for your verification only.
    Reject the question if answering it requires seeing any
    visual content that is not fully described in the text.

LEARNING MATERIAL:

{context}

QUESTION DATA:

{json.dumps(question_data, ensure_ascii=False)}

Return ONLY valid JSON:

{{
    "valid": true,
    "reason": "short verification explanation"
}}
"""

    verification_contents = [
        verification_prompt
    ]

    if relevant_sources:

        # Same PDF page images the generator saw
        # (max 2 per page).
        verification_contents.extend(
            collect_source_image_parts(
                relevant_sources,
                source_types=("pdf",),
                max_per_source=2
            )
        )

        for source in relevant_sources:

            if (
                source.get("source_type")
                == "slide"
            ):

                for image_data in source.get(
                    "images",
                    []
                )[:4]:

                    try:

                        verification_contents.append(
                            types.Part.from_bytes(
                                data=image_data[
                                    "bytes"
                                ],
                                mime_type=image_data[
                                    "mime_type"
                                ]
                            )
                        )

                    except Exception:
                        pass

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=verification_contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json"
        )
    )

    cleaned = clean_json_response(
        response.text
    )

    # An unreadable verdict counts as "not verified";
    # the generator then tries another question.
    try:
        verification = json.loads(
            cleaned
        )
    except ValueError:
        return False

    if not isinstance(verification, dict):
        return False

    return bool(
        verification.get(
            "valid",
            False
        )
    )

def generate_verified_assessment_question(
    source_scope,
    question_type,
    difficulty,
    topic_scope=""
):

    pool_items, pool_vectors = (
        get_assessment_pool(
            source_scope
        )
    )

    if not pool_items:

        raise ValueError(
            "No learning material is available "
            "for the selected source."
        )

    if "seen_questions" not in st.session_state:

        st.session_state.seen_questions = set()

    max_attempts = 8

    question_focuses = [
        "definition and core concept",
        "interpretation and meaning",
        "comparison with related concepts",
        "application or usage",
        "formula and parameter reasoning",
        "example-based reasoning",
        "conceptual relationship",
        "practical interpretation",
        "identify the correct property",
        "cause and effect reasoning",
        "distinguish between similar concepts",
        "scenario-based application",
        "parameter interpretation",
        "conceptual misconception check",
        "realistic use-case reasoning"
    ]

    for attempt in range(max_attempts):

        question_focus = random.choice(
            question_focuses
        )

        # ---------------------------------------
        # SELECT SOURCE MATERIAL
        # ---------------------------------------

        if topic_scope.strip():

            if question_type == "Numerical":

                search_query = (
                    f"{topic_scope} formula calculation "
                    f"numerical example values distance"
                )

                candidate_sources = (
                    retrieve_unified_sources(
                        search_query,
                        pool_items,
                        pool_vectors,
                        min(
                            10,
                            len(pool_items)
                        )
                    )
                )

                relevant_sources = candidate_sources[
                    :min(
                        6,
                        len(candidate_sources)
                    )
                ]

            else:

                candidate_sources = (
                    retrieve_unified_sources(
                        topic_scope,
                        pool_items,
                        pool_vectors,
                        min(
                            8,
                            len(pool_items)
                        )
                    )
                )

                sample_size = min(
                    4,
                    len(candidate_sources)
                )

                relevant_sources = random.sample(
                    candidate_sources,
                    sample_size
                )

        else:

            sample_size = min(
                4,
                len(pool_items)
            )

            indices = random.sample(
                range(len(pool_items)),
                sample_size
            )

            relevant_sources = [
                pool_items[index]
                for index in indices
            ]
        context_parts = []
        allowed_citations = []

        for number, source in enumerate(
            relevant_sources,
            start=1
        ):

            citation = source["citation"]

            allowed_citations.append(
                citation
            )

            context_parts.append(
                f"""
SOURCE {number}

Citation:
{citation}

Content:
{source["text"]}
"""
            )

        context = "\n".join(
            context_parts
        )

        previous_questions = []

        for item in st.session_state.get(
            "quiz_history",
            []
        )[-15:]:

            if (
                isinstance(item, dict)
                and item.get("question")
            ):
                previous_questions.append(
                    item["question"]
                )

        current_question = st.session_state.get(
            "quiz_question"
        )

        if current_question:

            current_text = current_question.get(
                "question",
                ""
            )

            if current_text:

                previous_questions.append(
                    current_text
                )

        # ---------------------------------------
        # QUESTION TYPE
        # ---------------------------------------

        if question_type == "MCQ":

            type_rules = """
Create exactly four answer options.

Only ONE option may be correct.

The JSON field "options" must contain exactly
four strings.

The "correct_answer" field must exactly match
one of those four options.
"""

        elif question_type == "Short Answer":

            type_rules = """
Create a question that requires a short written answer.

Do not create answer options.

The "options" field must be an empty list.

The "correct_answer" field should contain a concise
reference answer.
"""

        else:

            type_rules = """
        Create a numerical problem ONLY if the provided
        learning material contains enough information to
        derive the answer.

        IMPORTANT:

        - Prefer formulas, worked examples, distances,
        dimensions, counts, percentages, or other numerical
        relationships explicitly present in the material.
        - If the material contains a worked numerical example,
        transform it into a self-contained question.
        - Include ALL values required to solve the problem
        directly inside the question.
        - Never invent numbers that are not supported by
        the learning material.
        - The student must not need to open the original slide,
        PDF, or video to know the required values.
        - Do not create answer options.
        - The "options" field must be an empty list.
        - The "correct_answer" must contain the numerical
        answer and unit when applicable.

        If no valid numerical question can be created from
        the supplied material, return:

        {
            "cannot_generate": true
        }
        """
        prompt = f"""
You are LearnLens AI.

Create ONE adaptive assessment question using ONLY
the provided learning material.

Question type:
{question_type}

Difficulty:
{difficulty}

Student-selected topic or scope:
{topic_scope if topic_scope.strip() else "Any relevant topic"}

Question focus for this generation attempt:
{question_focus}

The new question MUST test a different aspect from the
previous questions whenever the learning material supports it.
Do not merely change wording.

{type_rules}

IMPORTANT RULES:

1. Use ONLY the provided learning material.
2. Never use outside knowledge.
3. The question must be educationally meaningful.
4. The answer must be clearly supported by the material.
5. Use one of the EXACT source citations provided below.
6. Do not repeat or closely paraphrase any previous question.
   Test a different concept, relationship, example, or reasoning skill.
7. Write in the same language as the learning material.
8. The question MUST be self-contained.
9. Never refer to hidden context using phrases such as
   "the example above", "the given example", "the figure",
   "the table", "the slide", or similar expressions unless
   all information needed to answer is explicitly included
   in the question text.
10. If the question is based on an example, calculation,
    observation, formula, or numerical data from the material,
    reproduce all necessary values and information inside
    the question itself.
11. The student must be able to answer the question without
    opening the original PDF, slide, or video.
12. Return ONLY valid JSON.
13. Do not use Markdown.
14. Any attached images are ONLY for your understanding of
    the material. The student will NOT see any image, chart,
    figure, equation image or table. Never ask about visual
    content; if a visual detail is needed, describe it fully
    in the question text, otherwise test a different concept.
Previously used questions:

{json.dumps(previous_questions, ensure_ascii=False)}

LEARNING MATERIAL:

{context}

Return this exact JSON structure:

{{
    "question": "question text",
    "question_type": "{question_type}",
    "options": [],
    "correct_answer": "correct answer",
    "explanation": "short source-grounded explanation",
    "topic": "short topic name",
    "difficulty": "{difficulty}",
    "source_citation": "exact citation from the material"
}}
"""

        # PDF figures of the selected pages help the model
        # understand the material (max 2 per page). The verifier
        # receives the same images.
        generation_contents = [prompt]

        generation_contents.extend(
            collect_source_image_parts(
                relevant_sources,
                source_types=("pdf",),
                max_per_source=2
            )
        )

        # API errors (including quota) propagate immediately so
        # a failing request is not retried 8 times.
        response = client.models.generate_content(
            model=GENERATION_MODEL,
            contents=generation_contents,
            config=types.GenerateContentConfig(
                response_mime_type="application/json"
            )
        )

        cleaned = clean_json_response(
            response.text
        )

        # A malformed response only skips this attempt.
        try:
            question_data = json.loads(
                cleaned
            )
        except ValueError:
            continue

        if not isinstance(question_data, dict):
            continue

        if question_data.get(
            "cannot_generate"
        ):

            continue

        question_text = str(
            question_data.get(
                "question",
                ""
            )
        ).strip()

        if not question_text:

            continue

        # Questions pointing at a figure / chart / slide the
        # student cannot see are rejected before verification.
        if references_hidden_visual(question_text):

            continue

        normalized_question = (
            normalize_question_text(
                question_text
            )
        )

        duplicate_found = False

        for old_question in previous_questions:

            if is_similar_question(
                question_text,
                old_question
            ):

                duplicate_found = True
                break

        if duplicate_found:
            continue

        question_data[
            "question_type"
        ] = question_type

        question_data[
            "difficulty"
        ] = difficulty

        source_citation = (
            question_data.get(
                "source_citation",
                ""
            )
        )

        if (
            source_citation
            not in allowed_citations
        ):

            continue

        if question_type == "MCQ":

            options = question_data.get(
                "options",
                []
            )

            if len(options) != 4:

                continue

            if (
                question_data.get(
                    "correct_answer"
                )
                not in options
            ):

                continue

        else:

            question_data[
                "options"
            ] = []

        is_valid = (
            verify_assessment_question(
                question_data,
                context,
                relevant_sources
            )
        )

        if not is_valid:

            continue

        st.session_state.seen_questions.add(
            normalized_question
        )

        # Canonical topic: the knowledge tag of the cited source
        # unit, so quiz results, mastery and tutor conversations
        # share topic names. Missing tags for these few units are
        # generated once (at most one batched call, cached).
        question_data["generated_topic"] = str(
            question_data.get("topic", "")
        ).strip()

        try:
            auto_tag_knowledge_items(relevant_sources)
        except Exception:
            pass

        cited_source = next(
            (
                source for source in relevant_sources
                if source.get("citation") == source_citation
            ),
            None
        )

        question_data["topic"] = canonical_topic_name(
            source_topic_tag(cited_source)
            or question_data["generated_topic"]
        )

        return question_data

    raise ValueError(
        "LearnLens could not generate a verified "
        "non-repeated question from the selected material. "
        "Try another topic, source, or question type."
    )


def grade_open_answer(
    question_data,
    student_answer
):

    prompt = f"""
You are LearnLens AI's assessment grader.

Evaluate the student's answer.

Question:
{question_data["question"]}

Reference answer:
{question_data["correct_answer"]}

Explanation:
{question_data["explanation"]}

Student answer:
{student_answer}

Rules:

1. Judge semantic correctness, not exact wording.
2. For numerical answers, accept equivalent numerical
   representations when mathematically correct.
3. Do not require information beyond the reference answer.
4. Return ONLY valid JSON.

Return:

{{
    "is_correct": true,
    "feedback": "short educational feedback"
}}
"""

    response = client.models.generate_content(
        model=GENERATION_MODEL,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json"
        )
    )

    cleaned = clean_json_response(
        response.text
    )

    try:
        result = json.loads(
            cleaned
        )
    except ValueError:
        raise ValueError(
            "The grader returned an unreadable response. "
            "Please submit your answer again."
        )

    if not isinstance(result, dict):
        raise ValueError(
            "The grader returned an unexpected response. "
            "Please submit your answer again."
        )

    return (
        bool(
            result.get(
                "is_correct",
                False
            )
        ),
        result.get(
            "feedback",
            ""
        )
    )


# ---------------------------------------------------
# LEARNER MASTERY MODEL (EMA)
# ---------------------------------------------------

APP_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

DEFAULT_TOPIC_MASTERY = 0.50

# Quiz evidence: correct = 1.0, incorrect = 0.0
QUIZ_MASTERY_ALPHA = 0.30

# Conversation-based heuristic evidence. Asking the tutor
# about a topic does not prove the student learned it, so
# these updates are deliberately small.
CONVERSATION_EVIDENCE = 0.55
CONVERSATION_ALPHA = 0.05

CONVERSATION_STRUGGLE_EVIDENCE = 0.35
CONVERSATION_STRUGGLE_ALPHA = 0.08

STRUGGLE_SIGNALS = [
    "anlamadım",
    "anlamadim",
    "anlamıyorum",
    "anlamiyorum",
    "zorlanıyorum",
    "zorlaniyorum",
    "karıştırıyorum",
    "karistiriyorum",
    "karıştırdım",
    "kafam karıştı",
    "bilmiyorum",
    "i don't understand",
    "i dont understand",
    "i do not understand",
    "i didn't understand",
    "i'm confused",
    "im confused",
    "i am confused",
    "i'm struggling",
    "i am struggling",
    "i don't know",
    "i dont know",
    "i mix up",
    "i keep mixing"
]

REFUSAL_MESSAGE = (
    "Bu bilgi yüklenen ders materyallerinde bulunamadı."
)

# Compared after normalize_for_matching(): lowercase, Turkish
# letters folded to ASCII, punctuation / Markdown removed.
REFUSAL_MARKERS = [
    "yuklenen ders materyallerinde bulunamadi",
    "ders materyallerinde bulunamadi",
    "yuklenen materyallerde bulunamadi",
    "not found in the uploaded",
    "not available in the uploaded",
    "not covered in the uploaded",
    "not contained in the uploaded"
]

# Phrases that point at visual content the student cannot see
# in the assessment interface (checked after normalization).
HIDDEN_VISUAL_PHRASES = [
    "sekildeki",
    "sekle gore",
    "sekilde gosteril",
    "sekilde goruldugu",
    "grafikteki",
    "grafige gore",
    "grafikte gosteril",
    "grafikte goruldugu",
    "gorseldeki",
    "gorsele gore",
    "gorselde",
    "resimdeki",
    "resme gore",
    "slayttaki",
    "slayta gore",
    "slaytta",
    "the figure",
    "figure above",
    "figure below",
    "the chart",
    "the graph",
    "the diagram",
    "the image",
    "the picture",
    "the slide",
    "shown above",
    "shown below"
]


def normalize_for_matching(text):
    """
    Deterministic normalization for phrase matching:
    Turkish dotless/dotted i, diacritics, case, Markdown,
    punctuation and whitespace differences are ignored.
    """

    text = str(text or "")

    text = text.replace("İ", "i").replace("I", "i")
    text = text.replace("ı", "i")

    text = unicodedata.normalize("NFKD", text)

    text = "".join(
        character for character in text
        if not unicodedata.combining(character)
    )

    text = text.casefold()

    text = re.sub(r"[^\w\s]", " ", text)
    text = text.replace("_", " ")

    return " ".join(text.split())


def references_hidden_visual(question_text):

    normalized = normalize_for_matching(question_text)

    return any(
        phrase in normalized
        for phrase in HIDDEN_VISUAL_PHRASES
    )


def describe_api_error(error):
    """
    Short, user-facing error text (no long API payloads).
    """

    if is_quota_error(error):
        return QUOTA_ERROR_MESSAGE

    message = " ".join(str(error).split())

    if len(message) > 300:
        message = message[:300] + "…"

    return message or type(error).__name__


def normalize_topic_key(topic):

    return " ".join(
        unicodedata.normalize(
            "NFC",
            str(topic or "")
        ).split()
    ).casefold()


def source_topic_tag(source):
    """
    Topic metadata of a knowledge unit (knowledge tag or video
    segment topic), or "" if it has none.
    """

    if not isinstance(source, dict):
        return ""

    topic = " ".join(
        str(source.get("topic", "")).split()
    )

    if topic == "General":
        return ""

    return topic


def canonical_topic_name(topic):
    """
    Maps a topic label to the spelling already used by the
    learner model or the knowledge tags when they differ only in
    whitespace / casing. Different topics are never merged.
    """

    cleaned = " ".join(
        unicodedata.normalize(
            "NFC",
            str(topic or "")
        ).split()
    )

    if not cleaned:
        return "General"

    key = normalize_topic_key(cleaned)

    known_topics = list(
        st.session_state.get("topic_performance", {})
    ) + list(
        st.session_state.get("learner_mastery", {})
    )

    cache = st.session_state.get("knowledge_tags_cache")

    if isinstance(cache, dict):
        known_topics += [
            tags.get("topic", "")
            for tags in cache.values()
            if isinstance(tags, dict)
        ]

    for known_topic in known_topics:

        if (
            known_topic
            and normalize_topic_key(known_topic) == key
        ):
            return known_topic

    return cleaned


def ema_update(old_mastery, evidence, alpha):
    """
    The single mastery update rule used everywhere
    (quiz answers, tutor conversations, simulations):

        new = old * (1 - alpha) + evidence * alpha
    """

    new_mastery = (
        old_mastery * (1 - alpha)
        + evidence * alpha
    )

    return max(0.0, min(1.0, new_mastery))


def update_learner_mastery(topic, evidence, alpha):

    if not isinstance(
        st.session_state.get("learner_mastery"),
        dict
    ):
        st.session_state.learner_mastery = {}

    topic = canonical_topic_name(topic)

    old_mastery = float(
        st.session_state.learner_mastery.get(
            topic,
            DEFAULT_TOPIC_MASTERY
        )
    )

    new_mastery = ema_update(
        old_mastery,
        evidence,
        alpha
    )

    st.session_state.learner_mastery[topic] = round(
        new_mastery,
        4
    )

    return new_mastery


def record_quiz_mastery(topic, is_correct):

    return update_learner_mastery(
        topic,
        1.0 if is_correct else 0.0,
        QUIZ_MASTERY_ALPHA
    )


def is_refusal_answer(answer):

    text = normalize_for_matching(answer)

    return any(
        marker in text
        for marker in REFUSAL_MARKERS
    )


def has_struggle_signal(question):

    text = unicodedata.normalize(
        "NFC",
        str(question or "")
    ).casefold()

    return any(
        signal in text
        for signal in STRUGGLE_SIGNALS
    )


def detect_conversation_topic(question, retrieved_sources):
    """
    1. Topic metadata of the most relevant retrieved source
       (knowledge tags or video segment topic).
    2. A topic already in the learner model that appears in
       the question or in the top retrieved text.
    3. Fallback: file name of the top retrieved source.
    """

    for source in retrieved_sources:

        topic = source_topic_tag(source)

        if topic:
            return topic

    known_topics = set(
        st.session_state.get("learner_mastery", {})
    ) | set(
        st.session_state.get("topic_performance", {})
    )

    haystack = str(question).casefold()

    if retrieved_sources:
        haystack += " " + str(
            retrieved_sources[0].get("text", "")
        ).casefold()

    for topic in sorted(
        known_topics,
        key=len,
        reverse=True
    ):

        if topic and topic.casefold() in haystack:
            return topic

    if retrieved_sources:

        file_stem = os.path.splitext(
            str(retrieved_sources[0].get("file", ""))
        )[0].strip()

        if file_stem:
            return file_stem

    return "General"


def record_conversation_mastery(
    question,
    retrieved_sources,
    answer
):
    """
    Conversation-based heuristic evidence.

    A successful tutoring interaction is weak evidence of
    progress (evidence 0.55, alpha 0.05). If the student
    explicitly signals difficulty ("anlamadım",
    "zorlanıyorum", "I'm confused" ...), it is weak evidence
    of a gap (evidence 0.35, alpha 0.08). Refusals
    (off-material questions) do not update the model.
    """

    if is_refusal_answer(answer):
        return None

    topic = canonical_topic_name(
        detect_conversation_topic(
            question,
            retrieved_sources
        )
    )

    struggle = has_struggle_signal(question)

    if struggle:
        evidence = CONVERSATION_STRUGGLE_EVIDENCE
        alpha = CONVERSATION_STRUGGLE_ALPHA
    else:
        evidence = CONVERSATION_EVIDENCE
        alpha = CONVERSATION_ALPHA

    new_mastery = update_learner_mastery(
        topic,
        evidence,
        alpha
    )

    if not isinstance(
        st.session_state.get("conversation_counts"),
        dict
    ):
        st.session_state.conversation_counts = {}

    st.session_state.conversation_counts[topic] = (
        int(
            st.session_state.conversation_counts.get(
                topic,
                0
            )
        )
        + 1
    )

    return {
        "topic": topic,
        "mastery": new_mastery,
        "struggle": struggle
    }


def mastery_status_label(mastery):

    if mastery >= 0.75:
        return "🟢 Mastered"

    if mastery >= 0.50:
        return "🟡 Developing"

    return "🔴 Needs Review"


# ---------------------------------------------------
# QUESTION REPETITION RATE
# ---------------------------------------------------

def compute_question_repetition_rate(quiz_history):
    """
    Compares every accepted question with all earlier
    accepted questions using the same is_similar_question
    rule (and threshold) as duplicate prevention.
    """

    questions = [
        str(item.get("question", "")).strip()
        for item in quiz_history
        if isinstance(item, dict)
        and str(item.get("question", "")).strip()
    ]

    repeated_questions = 0

    for index, question in enumerate(questions):

        for earlier_question in questions[:index]:

            if is_similar_question(
                question,
                earlier_question
            ):
                repeated_questions += 1
                break

    total_questions = len(questions)

    if total_questions:
        rate = repeated_questions / total_questions
    else:
        rate = 0.0

    return {
        "repeated_questions": repeated_questions,
        "total_questions": total_questions,
        "question_repetition_rate": rate
    }


def format_repetition_rate(repetition):

    return (
        f"{repetition['repeated_questions']} / "
        f"{repetition['total_questions']} repeated — "
        f"{repetition['question_repetition_rate'] * 100:.1f}%"
    )


# ---------------------------------------------------
# SIMULATED PERSONALIZATION BENCHMARK
# ---------------------------------------------------

PERSONALIZATION_PROFILES = [
    {
        "profile": "Beginner",
        "initial_mastery": 0.30,
        "session_1": [False, False, True],
        "session_2": [True, True, True]
    },
    {
        "profile": "Intermediate",
        "initial_mastery": 0.50,
        "session_1": [False, True, True],
        "session_2": [True, True, True]
    },
    {
        "profile": "Advanced",
        "initial_mastery": 0.70,
        "session_1": [True, True, True],
        "session_2": [True, True, True]
    }
]


def format_evidence_sequence(sequence):

    return ", ".join(
        "correct" if is_correct else "incorrect"
        for is_correct in sequence
    )


def run_personalization_benchmark():
    """
    Deterministic simulation (no API calls) that replays
    synthetic quiz evidence through the same ema_update rule
    used for real learners.
    """

    rows = []

    for profile in PERSONALIZATION_PROFILES:

        mastery = profile["initial_mastery"]

        for is_correct in profile["session_1"]:
            mastery = ema_update(
                mastery,
                1.0 if is_correct else 0.0,
                QUIZ_MASTERY_ALPHA
            )

        after_session_1 = mastery

        for is_correct in profile["session_2"]:
            mastery = ema_update(
                mastery,
                1.0 if is_correct else 0.0,
                QUIZ_MASTERY_ALPHA
            )

        after_session_2 = mastery

        rows.append({
            "Profile": profile["profile"],
            "Session 1 Evidence": format_evidence_sequence(
                profile["session_1"]
            ),
            "Session 2 Evidence": format_evidence_sequence(
                profile["session_2"]
            ),
            "Initial Mastery": round(
                profile["initial_mastery"],
                3
            ),
            "After Session 1": round(after_session_1, 3),
            "After Session 2": round(after_session_2, 3),
            "Mastery Gain": round(
                after_session_2 - profile["initial_mastery"],
                3
            )
        })

    return rows


# ---------------------------------------------------
# BATCH / BENCHMARK RAGAS EVALUATION
# ---------------------------------------------------

EVALUATION_TESTSET_FILE = os.path.join(
    APP_DIR,
    "evaluation_testset.json"
)

EVALUATION_RESULTS_JSON = os.path.join(
    APP_DIR,
    "evaluation_results.json"
)

EVALUATION_RESULTS_CSV = os.path.join(
    APP_DIR,
    "evaluation_results.csv"
)

RAGAS_METRIC_KEYS = [
    "faithfulness",
    "answer_relevancy",
    "context_precision",
    "context_recall"
]

DEFAULT_EVALUATION_TESTSET = {
    "name": "Kümeleme demo test set",
    "material": "Kümeleme.pptx",
    "test_cases": [
        {
            "id": "kumeleme_01",
            "question": "Minkowski uzaklığı nedir?",
            "reference_answer": (
                "Minkowski uzaklığı, p sayıda değişken "
                "dikkate alınarak gözlemler arasındaki "
                "uzaklığı hesaplamak için kullanılan genel "
                "bir uzaklık bağıntısıdır."
            ),
            "expected_source": "Kümeleme.pptx, Slide 7",
            "should_refuse": False
        },
        {
            "id": "kumeleme_02",
            "question": (
                "Simetrik uzaklık matrisinde d(i,j) ile "
                "d(j,i) arasındaki ilişki nedir?"
            ),
            "reference_answer": (
                "Simetrik uzaklık matrisinde d(i,j)=d(j,i) "
                "kabul edilir."
            ),
            "expected_source": "Kümeleme.pptx, Slide 4",
            "should_refuse": False
        },
        {
            "id": "kumeleme_03",
            "question": (
                "En yakın komşu algoritmasında iki küme "
                "arasındaki uzaklık nasıl belirlenir?"
            ),
            "reference_answer": (
                "İki kümenin içerdiği gözlemler arasında "
                "birbirine en yakın olanların uzaklığı iki "
                "küme arasındaki uzaklık olarak alınır."
            ),
            "expected_source": "Kümeleme.pptx, Slide 12",
            "should_refuse": False
        },
        {
            "id": "kumeleme_04",
            "question": (
                "En uzak komşu algoritmasında iki küme "
                "arasındaki uzaklık nasıl belirlenir?"
            ),
            "reference_answer": (
                "İki kümenin birbirine en uzak elemanları "
                "arasındaki mesafe kümeler arasındaki "
                "uzaklık olarak alınır."
            ),
            "expected_source": "Kümeleme.pptx, Slide 19",
            "should_refuse": False
        },
        {
            "id": "kumeleme_05",
            "question": (
                "k-ortalamalar yönteminin temel amacı nedir?"
            ),
            "reference_answer": (
                "k-ortalamalar yönteminin amacı verilen K "
                "değeri için toplam kare hatayı minimize "
                "eden kümeleri bulmaktır."
            ),
            "expected_source": "Kümeleme.pptx, Slide 24",
            "should_refuse": False
        },
        {
            "id": "off_material_01",
            "question": (
                "Yapay sinir ağlarında backpropagation "
                "algoritması nedir?"
            ),
            "reference_answer": REFUSAL_MESSAGE,
            "expected_source": "",
            "should_refuse": True
        }
    ]
}


def write_json_file(path, data):

    temporary_file = path + ".tmp"

    with open(
        temporary_file,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            data,
            file,
            ensure_ascii=False,
            indent=4
        )

    os.replace(temporary_file, path)


def load_evaluation_testset():
    """
    Returns the list of test cases. Creates the default
    test set if the file is missing; raises ValueError with
    a readable message if the file is malformed.
    """

    if not os.path.exists(EVALUATION_TESTSET_FILE):

        write_json_file(
            EVALUATION_TESTSET_FILE,
            DEFAULT_EVALUATION_TESTSET
        )

    try:

        with open(
            EVALUATION_TESTSET_FILE,
            "r",
            encoding="utf-8"
        ) as file:

            data = json.load(file)

    except json.JSONDecodeError as e:

        raise ValueError(
            "evaluation_testset.json is not valid JSON "
            f"(line {e.lineno}, column {e.colno}). "
            "Fix the file or delete it to restore the "
            "default test set."
        )

    if isinstance(data, dict):
        test_cases = data.get("test_cases")
    else:
        test_cases = data

    if not isinstance(test_cases, list) or not test_cases:

        raise ValueError(
            "evaluation_testset.json must contain a "
            "non-empty \"test_cases\" list."
        )

    cleaned_cases = []

    for number, case in enumerate(test_cases, start=1):

        if (
            not isinstance(case, dict)
            or not str(case.get("question", "")).strip()
        ):

            raise ValueError(
                f"Test case #{number} in "
                "evaluation_testset.json has no question."
            )

        should_refuse = bool(
            case.get("should_refuse", False)
        )

        reference_answer = str(
            case.get("reference_answer", "")
        ).strip()

        if not should_refuse and not reference_answer:

            raise ValueError(
                f"Test case #{number} in "
                "evaluation_testset.json needs a "
                "reference_answer."
            )

        cleaned_cases.append({
            "id": str(case.get("id", f"case_{number}")),
            "question": str(case["question"]).strip(),
            "reference_answer": reference_answer,
            "expected_source": str(
                case.get("expected_source", "")
            ).strip(),
            "should_refuse": should_refuse
        })

    return cleaned_cases


def filter_items_by_scope(items, vectors, scope):

    scope_types = {
        "PDF": "pdf",
        "Slides": "slide",
        "Video": "video"
    }

    filtered_items = []
    filtered_vectors = []

    for index, item in enumerate(items):

        if (
            scope == "All Materials"
            or item["source_type"] == scope_types.get(scope)
        ):

            filtered_items.append(item)
            filtered_vectors.append(vectors[index])

    if not filtered_items:
        return [], None

    return (
        filtered_items,
        np.array(filtered_vectors, dtype=np.float32)
    )


def normalize_citation_text(text):

    text = unicodedata.normalize(
        "NFC",
        str(text or "")
    ).casefold()

    text = text.replace("[", " ").replace("]", " ")
    text = re.sub(r"\s*,\s*", ", ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def expected_source_matches(
    expected_source,
    retrieved_sources,
    answer
):
    """
    Returns (in_retrieved, in_answer). "Slide 7" must not
    match "Slide 70", hence the digit look-ahead.
    """

    expected = normalize_citation_text(expected_source)

    if not expected:
        return False, False

    pattern = re.compile(
        re.escape(expected) + r"(?!\d)"
    )

    in_retrieved = any(
        pattern.search(
            normalize_citation_text(
                source.get("citation", "")
            )
        )
        for source in retrieved_sources
    )

    in_answer = bool(
        pattern.search(
            normalize_citation_text(answer)
        )
    )

    return in_retrieved, in_answer


def clean_metric_value(value):

    try:
        value = float(value)
    except (TypeError, ValueError):
        return None

    if math.isnan(value) or math.isinf(value):
        return None

    return round(value, 4)


def run_benchmark_case(
    case,
    items,
    vectors,
    learner_profile,
    topic_performance
):
    """
    Runs one benchmark case. Never raises: any failure is
    stored in the "error" field so the benchmark continues.
    """

    result = {
        "id": case["id"],
        "question": case["question"],
        "should_refuse": case["should_refuse"],
        "expected_source": case["expected_source"],
        "status": "error",
        "refused": None,
        "refusal_correct": None,
        "source_match": None,
        "source_in_retrieved": None,
        "source_in_answer": None,
        "faithfulness": None,
        "answer_relevancy": None,
        "context_precision": None,
        "context_recall": None,
        "retrieved_citations": [],
        "answer": "",
        "error": ""
    }

    try:

        retrieved_sources = retrieve_unified_sources(
            case["question"],
            items,
            vectors,
            min(6, len(items))
        )

        result["retrieved_citations"] = [
            source["citation"]
            for source in retrieved_sources
        ]

        if not retrieved_sources:
            raise ValueError(
                "Retrieval returned no sources for this question."
            )

        answer = answer_unified_question(
            case["question"],
            retrieved_sources,
            learner_profile,
            topic_performance
        )

        result["answer"] = answer

        refused = is_refusal_answer(answer)

        result["refused"] = refused
        result["refusal_correct"] = (
            refused == case["should_refuse"]
        )

        if case["expected_source"]:

            in_retrieved, in_answer = (
                expected_source_matches(
                    case["expected_source"],
                    retrieved_sources,
                    answer
                )
            )

            result["source_in_retrieved"] = in_retrieved
            result["source_in_answer"] = in_answer
            result["source_match"] = (
                in_retrieved or in_answer
            )

        if not case["should_refuse"]:

            scores = run_ragas_evaluation(
                case["question"],
                answer,
                [
                    source["text"]
                    for source in retrieved_sources
                ],
                case["reference_answer"]
            )

            for key in RAGAS_METRIC_KEYS:
                result[key] = clean_metric_value(
                    scores.get(key)
                )

        result["status"] = "ok"

    except Exception as e:

        result["error"] = str(e)[:300]

        if is_quota_error(e):
            result["status"] = "quota_error"

    return result


QUICK_BENCHMARK_IN_MATERIAL_CASES = 2
QUICK_BENCHMARK_REFUSAL_CASES = 1
BENCHMARK_CASE_DELAY_SECONDS = 5

QUOTA_ERROR_MESSAGE = (
    "Gemini API quota reached. Wait for quota reset and retry."
)


def is_quota_error(error):
    """
    Detects Gemini 429 / RESOURCE_EXHAUSTED errors, also when
    RAGAS or the OpenAI client wraps the original exception.
    """

    current = error
    checked = 0

    while current is not None and checked < 5:

        status_code = getattr(current, "status_code", None)
        code = getattr(current, "code", None)

        if status_code == 429 or code == 429:
            return True

        text = str(current).upper()

        if (
            "RESOURCE_EXHAUSTED" in text
            or re.search(r"\b429\b", text)
        ):
            return True

        current = (
            current.__cause__
            or current.__context__
        )
        checked += 1

    return False


def select_quick_benchmark_cases(test_cases):
    """
    Quota-safe subset: first off-material (refusal) case +
    first 2 in-material cases. The refusal case runs first
    because it needs no RAGAS calls (cheapest), so it still
    completes if the quota runs out later.
    """

    in_material = [
        case for case in test_cases
        if not case["should_refuse"]
    ][:QUICK_BENCHMARK_IN_MATERIAL_CASES]

    refusal = [
        case for case in test_cases
        if case["should_refuse"]
    ][:QUICK_BENCHMARK_REFUSAL_CASES]

    return refusal + in_material


def skipped_benchmark_case(case):

    return {
        "id": case["id"],
        "question": case["question"],
        "should_refuse": case["should_refuse"],
        "expected_source": case["expected_source"],
        "status": "skipped",
        "refused": None,
        "refusal_correct": None,
        "source_match": None,
        "source_in_retrieved": None,
        "source_in_answer": None,
        "faithfulness": None,
        "answer_relevancy": None,
        "context_precision": None,
        "context_recall": None,
        "retrieved_citations": [],
        "answer": "",
        "error": "Not run: Gemini API quota reached."
    }


def summarize_benchmark(results):

    summary = {}

    # Only fully successful cases count towards scores, so a
    # quota-interrupted run is never shown as a valid result.
    successful_results = [
        result for result in results
        if result.get("status") == "ok"
    ]

    for key in RAGAS_METRIC_KEYS:

        values = [
            result[key]
            for result in successful_results
            if result.get(key) is not None
        ]

        summary[f"average_{key}"] = (
            round(sum(values) / len(values), 4)
            if values
            else None
        )

        summary[f"{key}_cases"] = len(values)

    refusal_checks = [
        result["refusal_correct"]
        for result in successful_results
        if result.get("refusal_correct") is not None
    ]

    summary["refusal_correct"] = sum(
        1 for value in refusal_checks if value
    )
    summary["refusal_total"] = len(refusal_checks)
    summary["refusal_accuracy"] = (
        round(
            summary["refusal_correct"]
            / len(refusal_checks),
            4
        )
        if refusal_checks
        else None
    )

    source_checks = [
        result["source_match"]
        for result in successful_results
        if result.get("source_match") is not None
    ]

    summary["source_match_correct"] = sum(
        1 for value in source_checks if value
    )
    summary["source_match_total"] = len(source_checks)
    summary["source_match_rate"] = (
        round(
            summary["source_match_correct"]
            / len(source_checks),
            4
        )
        if source_checks
        else None
    )

    summary["total_cases"] = len(results)
    summary["successful_cases"] = len(successful_results)
    summary["failed_cases"] = sum(
        1 for result in results
        if result.get("status") in ("error", "quota_error")
    )
    summary["skipped_cases"] = sum(
        1 for result in results
        if result.get("status") == "skipped"
    )
    summary["quota_exhausted"] = any(
        result.get("status") == "quota_error"
        for result in results
    )

    if summary["quota_exhausted"]:
        summary["run_status"] = (
            "stopped_quota_partial"
            if successful_results
            else "stopped_quota"
        )
    elif not successful_results:
        summary["run_status"] = "failed"
    elif summary["failed_cases"]:
        summary["run_status"] = "completed_with_errors"
    else:
        summary["run_status"] = "completed"

    return summary


def save_benchmark_results(
    results,
    summary,
    scope,
    benchmark_mode="Quick Benchmark"
):

    report = {
        "generated_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "source_scope": scope,
        "benchmark_mode": benchmark_mode,
        "generation_model": GENERATION_MODEL,
        "run_status": summary.get("run_status"),
        "successful_cases": summary.get("successful_cases"),
        "failed_cases": summary.get("failed_cases"),
        "quota_exhausted": summary.get("quota_exhausted"),
        "summary": summary,
        "cases": results
    }

    write_json_file(EVALUATION_RESULTS_JSON, report)

    # Run-level fields repeated on every CSV row.
    run_fields = {
        "run_status": summary.get("run_status"),
        "successful_cases": summary.get("successful_cases"),
        "failed_cases": summary.get("failed_cases"),
        "quota_exhausted": summary.get("quota_exhausted")
    }

    csv_columns = [
        "run_status",
        "successful_cases",
        "failed_cases",
        "quota_exhausted",
        "id",
        "question",
        "should_refuse",
        "status",
        "refused",
        "refusal_correct",
        "expected_source",
        "source_match",
        "source_in_retrieved",
        "source_in_answer",
        "faithfulness",
        "answer_relevancy",
        "context_precision",
        "context_recall",
        "retrieved_citations",
        "error"
    ]

    # utf-8-sig keeps Turkish characters readable in Excel.
    with open(
        EVALUATION_RESULTS_CSV,
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=csv_columns
        )

        writer.writeheader()

        for result in results:

            row = {
                column: result.get(column)
                for column in csv_columns
            }

            row.update(run_fields)

            row["retrieved_citations"] = "; ".join(
                result.get("retrieved_citations", [])
            )

            writer.writerow(row)

    return report


def load_saved_benchmark_report():

    if not os.path.exists(EVALUATION_RESULTS_JSON):
        return None

    try:

        with open(
            EVALUATION_RESULTS_JSON,
            "r",
            encoding="utf-8"
        ) as file:

            report = json.load(file)

        if (
            isinstance(report, dict)
            and isinstance(report.get("cases"), list)
            and isinstance(report.get("summary"), dict)
        ):
            # Ignore malformed rows from a hand-edited file.
            report["cases"] = [
                case for case in report["cases"]
                if isinstance(case, dict)
            ]
            return report

    except Exception:
        pass

    return None


def format_metric(value, as_percent=False):

    if value is None:
        return "—"

    if as_percent:
        return f"{value * 100:.0f}%"

    return f"{value:.3f}"


# ---------------------------------------------------
# KNOWLEDGE UNIT TAGGING
# (topic / subtopic / key concepts / prerequisites)
# ---------------------------------------------------

KNOWLEDGE_TAGS_FILE = os.path.join(
    APP_DIR,
    "knowledge_tags.json"
)

KNOWLEDGE_TAG_BATCH_SIZE = 8


def knowledge_item_key(item):
    """
    Cache key = hash of citation + text, so the same
    material re-uploaded later reuses its tags.
    """

    base = (
        str(item.get("citation", ""))
        + "\n"
        + str(item.get("text", ""))
    )

    return hashlib.sha256(
        base.encode("utf-8")
    ).hexdigest()


def load_knowledge_tags():

    if not os.path.exists(KNOWLEDGE_TAGS_FILE):

        try:
            write_json_file(KNOWLEDGE_TAGS_FILE, {})
        except Exception:
            pass

        return {}

    try:

        with open(
            KNOWLEDGE_TAGS_FILE,
            "r",
            encoding="utf-8"
        ) as file:

            data = json.load(file)

        if isinstance(data, dict):
            return data

    except Exception:
        pass

    return {}


def get_knowledge_tags_cache():

    if not isinstance(
        st.session_state.get("knowledge_tags_cache"),
        dict
    ):
        st.session_state.knowledge_tags_cache = (
            load_knowledge_tags()
        )

    return st.session_state.knowledge_tags_cache


def as_string_list(value, limit=6):

    if isinstance(value, str):
        value = value.split(",")

    if not isinstance(value, list):
        return []

    cleaned = []

    for entry in value:

        entry = str(entry).strip()

        if entry and entry not in cleaned:
            cleaned.append(entry)

    return cleaned[:limit]


def apply_knowledge_tags(items):
    """
    Adds topic / subtopic / key_concepts / prerequisites
    from the local cache. Existing values (e.g. video
    segment topics) are kept.
    """

    cache = get_knowledge_tags_cache()

    for item in items:

        tags = cache.get(knowledge_item_key(item))

        if isinstance(tags, dict):

            current_topic = str(
                item.get("topic", "")
            ).strip()

            if not current_topic or current_topic == "General":
                item["topic"] = tags.get("topic", "")

            item["subtopic"] = (
                item.get("subtopic")
                or tags.get("subtopic", "")
            )

            item["key_concepts"] = (
                item.get("key_concepts")
                or list(tags.get("key_concepts", []))
            )

            item["prerequisites"] = (
                item.get("prerequisites")
                or list(tags.get("prerequisites", []))
            )

            item["tagged"] = True

        else:

            item.setdefault("subtopic", "")
            item.setdefault("key_concepts", [])
            item.setdefault("prerequisites", [])
            item["tagged"] = False

    return items


def tag_knowledge_items(items, progress_callback=None):
    """
    Tags untagged knowledge units with Gemini in batches of
    KNOWLEDGE_TAG_BATCH_SIZE (one API call per batch, never
    one per chunk). Results are cached in knowledge_tags.json.
    A failed batch is skipped; RAG keeps working without tags.
    """

    cache = get_knowledge_tags_cache()

    pending = []
    pending_keys = set()

    for item in items:

        key = knowledge_item_key(item)

        if key in cache or key in pending_keys:
            continue

        pending_keys.add(key)
        pending.append((key, item))

    tagged_count = 0
    failed_count = 0

    for start in range(
        0,
        len(pending),
        KNOWLEDGE_TAG_BATCH_SIZE
    ):

        # Topic names already in use, so batches reuse the same
        # wording (canonical topics without extra calls).
        existing_topic_names = sorted({
            str(tags.get("topic", "")).strip()
            for tags in cache.values()
            if isinstance(tags, dict)
            and str(tags.get("topic", "")).strip()
        })[:40]

        batch = pending[
            start:start + KNOWLEDGE_TAG_BATCH_SIZE
        ]

        batch_payload = []

        for number, (key, item) in enumerate(
            batch,
            start=1
        ):

            batch_payload.append({
                "id": f"item_{number}",
                "citation": item.get("citation", ""),
                "existing_topic": (
                    item.get("topic", "")
                    if item.get("source_type") == "video"
                    else ""
                ),
                "text": str(item.get("text", ""))[:1500]
            })

        prompt = f"""
You are organizing course material into a knowledge
structure for an adaptive tutoring system.

For EACH knowledge unit below, return:
- topic: the main course topic (short)
- subtopic: the more specific subtopic (short)
- key_concepts: 2-5 key concepts explicitly covered
- prerequisites: 0-3 concepts a student should know first

RULES:
1. Base tags only on the unit's text. Do not invent content.
2. If "existing_topic" is given, use it as the topic.
3. Use the same language as the material.
4. Keep every value short (a few words).
5. Return one entry per id. Return ONLY valid JSON.
6. Topics are broad course sections. If one of these existing
   topic names fits, reuse it with exactly the same spelling:
   {json.dumps(existing_topic_names, ensure_ascii=False)}

Return exactly this structure:

{{
    "items": [
        {{
            "id": "item_1",
            "topic": "topic",
            "subtopic": "subtopic",
            "key_concepts": ["concept"],
            "prerequisites": ["prerequisite"]
        }}
    ]
}}

KNOWLEDGE UNITS:

{json.dumps(batch_payload, ensure_ascii=False)}
"""

        try:

            response = client.models.generate_content(
                model=GENERATION_MODEL,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json"
                )
            )

            data = json.loads(
                clean_json_response(response.text)
            )

            if isinstance(data, dict):
                tag_results = data.get("items", [])
            else:
                tag_results = data

            results_by_id = {
                str(entry.get("id")): entry
                for entry in tag_results
                if isinstance(entry, dict)
            }

            for number, (key, item) in enumerate(
                batch,
                start=1
            ):

                entry = results_by_id.get(
                    f"item_{number}"
                )

                if not entry:
                    failed_count += 1
                    continue

                cache[key] = {
                    "citation": item.get("citation", ""),
                    "topic": str(
                        entry.get("topic", "")
                    ).strip(),
                    "subtopic": str(
                        entry.get("subtopic", "")
                    ).strip(),
                    "key_concepts": as_string_list(
                        entry.get("key_concepts", []),
                        limit=5
                    ),
                    "prerequisites": as_string_list(
                        entry.get("prerequisites", []),
                        limit=3
                    )
                }

                tagged_count += 1

            write_json_file(KNOWLEDGE_TAGS_FILE, cache)

        except Exception:

            failed_count += len(batch)

        if progress_callback:

            progress_callback(
                min(
                    start + KNOWLEDGE_TAG_BATCH_SIZE,
                    len(pending)
                ),
                len(pending)
            )

    return tagged_count, failed_count, len(pending)


def auto_tag_knowledge_items(items):
    """
    Generates tags only for the given units that are not cached
    yet (e.g. the sources of one tutor answer or one assessment
    question: at most one batched Gemini call). Units that failed
    once in this session are not retried automatically, so a
    failure never causes repeated calls on later reruns.
    Never raises; RAG works without tags.
    """

    if not st.session_state.get(
        "auto_tagging_enabled",
        True
    ):
        return 0, 0

    try:

        cache = get_knowledge_tags_cache()

        failed_keys = st.session_state.setdefault(
            "knowledge_tag_failed_keys",
            []
        )

        missing_items = []
        missing_keys = set()

        for item in items:

            key = knowledge_item_key(item)

            if (
                key in cache
                or key in failed_keys
                or key in missing_keys
            ):
                continue

            missing_keys.add(key)
            missing_items.append(item)

        if not missing_items:
            apply_knowledge_tags(items)
            return 0, 0

        tagged_count, failed_count, _ = tag_knowledge_items(
            missing_items
        )

        for key in missing_keys:
            if key not in cache:
                failed_keys.append(key)

        apply_knowledge_tags(items)

        return tagged_count, failed_count

    except Exception:

        return 0, 0


# ---------------------------------------------------
# TABS
# ---------------------------------------------------

# Tabs are displayed in demo order; each "with tabN:" block
# below keeps its original variable name.
tab8, tab1, tab6, tab5, tab2, tab7, tab3, tab4, tab9 = st.tabs([
    "🧭 Learner Profile",
    "📄 PDF RAG",
    "📽️ Slide Understanding",
    "🎞️ Video Understanding",
    "🖼️ Image Understanding",
    "🎓 Unified Tutor",
    "🧠 Adaptive Assessment",
    "📊 Mastery Dashboard",
    "🧪 Evaluation"
])
# ===================================================
# TAB 1 - PDF RAG
# ===================================================

with tab1:

    section_header(
        "📄 PDF Learning Materials",
        "Upload textbooks, notes or lecture PDFs. LearnLens indexes "
        "the text and relevant figures for cited answers.",
        eyebrow="Step 2 · Learning materials"
    )

    with st.container(key="card_upload_pdf"):

        uploaded_files = st.file_uploader(
            "Upload PDF learning materials",
            type=["pdf"],
            accept_multiple_files=True,
            key="pdf_uploader"
        )

    if uploaded_files:

        uploaded_data = [
            (uploaded_file.name, uploaded_file.getvalue())
            for uploaded_file in uploaded_files
        ]

        signature = create_file_signature(uploaded_data)

        pdf_failure = st.session_state.get("pdf_index_failure")

        if (
            pdf_failure
            and pdf_failure.get("signature") == signature
        ):

            # Do not call the embedding API again on every
            # rerun after a failure; retry only on request.
            st.error(
                "The PDF index could not be created: "
                + pdf_failure.get("message", "")
            )

            if st.button("Retry PDF processing", key="retry_pdf_index"):
                st.session_state.pop("pdf_index_failure", None)
                st.rerun()

        elif (
            "document_signature" not in st.session_state
            or st.session_state.document_signature != signature
        ):

            with st.spinner("Indexing learning materials…"):
                try:
                    chunks, full_text, file_infos = extract_pdf_chunks(uploaded_data)

                    if not chunks:
                        raise ValueError(
                            "No readable text or figures were "
                            "found in the PDFs."
                        )

                    embeddings = create_embeddings(chunks)

                    st.session_state.document_signature = signature
                    st.session_state.chunks = chunks
                    st.session_state.embeddings = embeddings
                    st.session_state.full_text = full_text
                    st.session_state.file_infos = file_infos

                except Exception as e:

                    for stale_key in (
                        "document_signature",
                        "chunks",
                        "embeddings",
                        "full_text",
                        "file_infos"
                    ):
                        st.session_state.pop(stale_key, None)

                    st.session_state.pdf_index_failure = {
                        "signature": signature,
                        "message": describe_api_error(e)
                    }

                    st.rerun()

        pdf_index_ready = (
            st.session_state.get("document_signature") == signature
            and "chunks" in st.session_state
        )

        if pdf_index_ready:

            chunks = st.session_state.chunks
            embeddings = st.session_state.embeddings
            full_text = st.session_state.full_text
            file_infos = st.session_state.file_infos

            pdf_figure_pages = {
                (chunk["file"], chunk["page"]): len(chunk.get("images", []))
                for chunk in chunks
            }

            pdf_figure_count = sum(pdf_figure_pages.values())

            status_badges([
                f"{len(file_infos)} file(s) indexed",
                f"{sum(info['pages'] for info in file_infos)} pages",
                f"{len(chunks)} knowledge chunks",
                f"{pdf_figure_count} figure(s)",
                f"Embeddings ready ({embeddings.shape[1]}-d)"
            ])

            with st.expander("📑 Indexed files and extracted text"):

                for info in file_infos:
                    st.write(
                        f"• {info['name']} — {info['pages']} page(s)"
                    )

                if pdf_figure_count:
                    st.caption(
                        f"🖼️ {pdf_figure_count} embedded figure(s) "
                        "extracted (max 2 per page); they are sent to "
                        "Gemini when their page is retrieved."
                    )

                st.text_area(
                    "PDF Content",
                    full_text,
                    height=300
                )

            st.divider()

            st.subheader("💬 Ask About Your PDFs")

            question = st.text_input(
                "Ask a question about the uploaded PDF material:",
                key="pdf_question"
            )

            if st.button("Ask PDF", type="primary"):

                if not question.strip():
                    st.warning("Please enter a question.")

                else:
                    try:
                        with st.spinner("Searching learning materials…"):
                            retrieved_chunks = retrieve_chunks(
                                question,
                                chunks,
                                embeddings,
                                TOP_K
                            )

                        context_parts = []

                        for number, chunk in enumerate(retrieved_chunks, start=1):
                            context_parts.append(
                                f"""
SOURCE {number}

File: {chunk['file']}
Page: {chunk['page']}

{chunk['text']}
"""
                            )

                        context = "\n".join(context_parts)

                        prompt = f"""
You are LearnLens AI, a personalized learning assistant.

The student asked a question about uploaded learning
materials.

You have been given only the most relevant sections
retrieved by a semantic search system.

IMPORTANT RULES:

1. Answer ONLY using the retrieved sources below.
2. Do not add outside information.
3. If the answer cannot be determined from these sources,
   clearly say that the information was not found in the
   uploaded learning material.
4. Explain the answer clearly as a tutor.
5. Answer in the same language as the student's question.
6. Cite the source at the relevant part of the answer using:

   [File name, Page X]

7. Do not invent page numbers or sources.

RETRIEVED SOURCES:

{context}

STUDENT QUESTION:

{question}
"""

                        # Figures from the retrieved pages
                        # (max 2 per page, each page once).
                        pdf_contents = [prompt]

                        pdf_contents.extend(
                            collect_source_image_parts(
                                retrieved_chunks,
                                source_types=("pdf",),
                                max_per_source=2
                            )
                        )

                        with st.spinner("Generating grounded answer…"):
                            response = client.models.generate_content(
                                model=GENERATION_MODEL,
                                contents=pdf_contents
                            )

                        with st.container(key="card_pdf_answer"):

                            eyebrow_label("Answer")

                            st.write(response.text)

                        eyebrow_label("Sources used")

                        st.caption(
                            " · ".join(
                                f"[{chunk['file']}, Page {chunk['page']}]"
                                for chunk in retrieved_chunks
                            )
                        )

                        with st.expander("📚 Retrieved passages"):
                            for number, chunk in enumerate(retrieved_chunks, start=1):
                                st.markdown(
                                    f"**{number}. {chunk['file']} — "
                                    f"Page {chunk['page']}**"
                                )
                                st.caption(
                                    f"Similarity score: {chunk['score']:.3f}"
                                )
                                st.write(chunk["text"])
                                st.divider()

                    except Exception as e:
                        st.error(f"An error occurred: {describe_api_error(e)}")

    else:

        st.info(
            "No PDF uploaded yet. Drop one or more PDFs above — "
            "they join the shared knowledge base used by the "
            "Unified Tutor and Adaptive Assessment."
        )

        # Uploader emptied: remove the PDF part of the
        # knowledge base so stale material is not used.
        for stale_key in (
            "document_signature",
            "chunks",
            "embeddings",
            "full_text",
            "file_infos",
            "pdf_index_failure"
        ):
            st.session_state.pop(stale_key, None)


# ===================================================
# TAB 2 - IMAGE UNDERSTANDING
# ===================================================

with tab2:

    section_header(
        "🖼️ Image Understanding",
        "Upload a diagram, chart, handwritten note or screenshot "
        "and ask LearnLens about it.",
        eyebrow="Learning materials"
    )

    with st.container(key="card_upload_image"):

        uploaded_image = st.file_uploader(
            "Upload an image",
            type=["png", "jpg", "jpeg", "webp"],
            accept_multiple_files=False,
            key="image_uploader"
        )

    if uploaded_image:

        image_col, ask_col = st.columns([1, 1], gap="large")

        with image_col:

            st.image(
                uploaded_image,
                caption=uploaded_image.name,
                width="stretch"
            )

        with ask_col:

            st.subheader("💬 Ask About This Image")

            image_question = st.text_input(
                "Ask a question about the uploaded image:",
                key="image_question"
            )

            if st.button("Ask Image", type="primary"):

                if not image_question.strip():
                    st.warning("Please enter a question.")

                else:
                    try:
                        with st.spinner("Analyzing image…"):
                            image_answer = answer_image_question(
                                uploaded_image,
                                image_question
                            )

                        with st.container(key="card_image_answer"):

                            eyebrow_label("Answer")

                            st.write(image_answer)

                    except Exception as e:
                        st.error(f"An error occurred while analyzing the image: {describe_api_error(e)}")

    else:

        st.info(
            "No image uploaded yet. Images are answered directly "
            "(they are not added to the shared knowledge base)."
        )

# ===================================================
# TAB 3 - ADAPTIVE ASSESSMENT
# ===================================================

with tab3:

    section_header(
        "🧠 Adaptive Assessment",
        "Verified, source-grounded questions from your learning "
        "materials. Difficulty adapts to your answers.",
        eyebrow="Step 5 · Practice"
    )

    unified_items, unified_vectors = (
        build_unified_knowledge_base()
    )

    if not unified_items:

        st.info(
            "No learning materials yet. Upload a PDF, presentation "
            "or video first; questions are generated only from "
            "your own learning materials."
        )

    else:

        # -------------------------------------------
        # INITIALIZE NEW SESSION VALUES
        # -------------------------------------------

        if "quiz_score" not in st.session_state:

            reset_quiz()

        if "seen_questions" not in st.session_state:
            st.session_state.seen_questions = set()

        if "quiz_history" not in st.session_state:
            st.session_state.quiz_history = []

        if "last_answer_correct" not in st.session_state:

            st.session_state.last_answer_correct = False

        if "last_feedback" not in st.session_state:

            st.session_state.last_feedback = ""

        # Eski quiz formatından kalmış soru varsa temizle
        if (
            st.session_state.quiz_question is not None
            and "question_type"
            not in st.session_state.quiz_question
        ):

            st.session_state.quiz_question = None
            st.session_state.quiz_answered = False

        # -------------------------------------------
        # ASSESSMENT SETTINGS
        # -------------------------------------------

        st.subheader(
            "⚙️ Assessment Settings"
        )

        available_types = {
            item["source_type"]
            for item in unified_items
        }

        source_options = [
            "All Materials"
        ]

        if "pdf" in available_types:

            source_options.append(
                "PDF"
            )

        if "slide" in available_types:

            source_options.append(
                "Slides"
            )

        if "video" in available_types:

            source_options.append(
                "Video"
            )

        col1, col2 = st.columns(2)

        with col1:

            source_scope = st.selectbox(
                "Assessment source",
                source_options,
                key="assessment_source"
            )

        with col2:

            question_type = st.selectbox(
                "Question type",
                [
                    "MCQ",
                    "Short Answer",
                    "Numerical"
                ],
                key="assessment_type"
            )

        topic_scope = st.text_input(
            "Topic or scope (optional)",
            placeholder=(
                "Example: Pandas, clustering, "
                "Minkowski distance..."
            ),
            key="assessment_topic"
        )

        st.caption(
            "Question difficulty adapts automatically: 2 correct "
            "answers in a row → harder, 1 wrong answer → easier. "
            "Every question is checked by a second verifier pass "
            "and cites its source."
        )

        st.divider()

        # -------------------------------------------
        # METRICS
        # -------------------------------------------

        accuracy = 0

        if (
            st.session_state.quiz_total
            > 0
        ):

            accuracy = (
                st.session_state.quiz_score
                / st.session_state.quiz_total
            ) * 100

        metric1, metric2, metric3, metric4 = (
            st.columns(4)
        )

        with metric1:

            st.metric(
                "Next Difficulty",
                st.session_state.quiz_difficulty
            )

        with metric2:

            st.metric(
                "Question",
                f"#{st.session_state.quiz_question_number}"
                if st.session_state.quiz_question_number
                else "—"
            )

        with metric3:

            st.metric(
                "Score",
                f"{st.session_state.quiz_score}/"
                f"{st.session_state.quiz_total}"
            )

        with metric4:

            st.metric(
                "Accuracy",
                f"{accuracy:.0f}%"
            )

        st.divider()

        # -------------------------------------------
        # START
        # -------------------------------------------

        if (
            st.session_state.quiz_question
            is None
        ):

            if st.button(
                "Start Assessment",
                type="primary"
            ):

                try:

                    with st.spinner(
                        "Generating adaptive question…"
                    ):

                        question_data = (
                            generate_verified_assessment_question(
                                source_scope,
                                question_type,
                                st.session_state.quiz_difficulty,
                                topic_scope
                            )
                        )

                        st.session_state.quiz_question = (
                            question_data
                        )

                        st.session_state.quiz_question_number += 1

                        st.session_state.quiz_answered = False

                    st.rerun()

                except Exception as e:

                    st.error(
                        f"Assessment question could not "
                        f"be created: {describe_api_error(e)}"
                    )

        else:

            question_data = (
                st.session_state.quiz_question
            )

            question_card = st.container(key="card_question")

            with question_card:

                with st.container(horizontal=True, gap="small"):

                    st.badge(
                        f"Question {st.session_state.quiz_question_number}",
                        color="gray"
                    )

                    difficulty_badge(
                        question_data["difficulty"]
                    )

                    st.badge(
                        question_data["question_type"],
                        color="blue"
                    )

                st.caption(
                    f"Topic: {question_data['topic']}"
                )

                st.markdown(
                    "#### " + str(question_data["question"])
                )

            student_answer = None

            # ---------------------------------------
            # MCQ
            # ---------------------------------------

            if (
                question_data[
                    "question_type"
                ]
                == "MCQ"
            ):

                student_answer = question_card.radio(
                    "Choose your answer:",
                    question_data[
                        "options"
                    ],
                    index=None,
                    key=(
                        "assessment_answer_"
                        + str(
                            st.session_state.quiz_question_number
                        )
                    ),
                    disabled=(
                        st.session_state.quiz_answered
                    )
                )

            # ---------------------------------------
            # SHORT ANSWER / NUMERICAL
            # ---------------------------------------

            else:

                student_answer = question_card.text_input(
                    "Your answer:",
                    key=(
                        "assessment_text_answer_"
                        + str(
                            st.session_state.quiz_question_number
                        )
                    ),
                    disabled=(
                        st.session_state.quiz_answered
                    )
                )

            # ---------------------------------------
            # SUBMIT ANSWER
            # ---------------------------------------

            if not st.session_state.quiz_answered:

                if st.button(
                    "Submit Answer",
                    type="primary"
                ):

                    if (
                        student_answer is None
                        or not str(
                            student_answer
                        ).strip()
                    ):

                        st.warning(
                            "Please enter or select an answer."
                        )

                    else:

                        try:

                            if (
                                question_data[
                                    "question_type"
                                ]
                                == "MCQ"
                            ):

                                is_correct = (
                                    student_answer
                                    == question_data[
                                        "correct_answer"
                                    ]
                                )

                                feedback = (
                                    question_data[
                                        "explanation"
                                    ]
                                )

                            else:

                                with st.spinner(
                                    "Checking your answer…"
                                ):

                                    (
                                        is_correct,
                                        feedback
                                    ) = grade_open_answer(
                                        question_data,
                                        student_answer
                                    )

                            st.session_state.quiz_total += 1

                            if is_correct:

                                st.session_state.quiz_score += 1

                            update_topic_performance(
                                question_data[
                                    "topic"
                                ],
                                is_correct
                            )

                            record_quiz_mastery(
                                question_data[
                                    "topic"
                                ],
                                is_correct
                            )

                            update_difficulty(
                                is_correct
                            )

                            st.session_state.quiz_answered = True

                            st.session_state.last_answer_correct = (
                                is_correct
                            )

                            st.session_state.last_feedback = (
                                feedback
                            )

                            st.session_state.quiz_history.append({
                                "question": (
                                    question_data[
                                        "question"
                                    ]
                                ),
                                "type": (
                                    question_data[
                                        "question_type"
                                    ]
                                ),
                                "difficulty": (
                                    question_data[
                                        "difficulty"
                                    ]
                                ),
                                "topic": (
                                    question_data[
                                        "topic"
                                    ]
                                ),
                                "source": (
                                    question_data[
                                        "source_citation"
                                    ]
                                ),
                                "student_answer": (
                                    student_answer
                                ),
                                "correct_answer": (
                                    question_data[
                                        "correct_answer"
                                    ]
                                ),
                                "is_correct": (
                                    is_correct
                                ),
                                "explanation": (
                                    question_data.get(
                                        "explanation",
                                        ""
                                    )
                                )
                            })

                            save_learner_state()

                            st.rerun()

                        except Exception as e:

                            st.error(
                                f"Answer evaluation failed: {describe_api_error(e)}"
                            )

            # ---------------------------------------
            # RESULT
            # ---------------------------------------

            else:

                if (
                    st.session_state.last_answer_correct
                ):

                    st.success(
                        "✓ Correct",
                        icon=":material/check_circle:"
                    )

                else:

                    st.error(
                        "✗ Needs Review",
                        icon=":material/error:"
                    )

                    st.write(
                        "**Correct answer:** "
                        + str(
                            question_data[
                                "correct_answer"
                            ]
                        )
                    )

                st.info(
                    "💡 "
                    + st.session_state.last_feedback
                )

                st.caption(
                    "📚 Source: "
                    + question_data[
                        "source_citation"
                    ]
                )

                st.caption(
                    "Next question difficulty: "
                    + st.session_state.quiz_difficulty
                )

                # -----------------------------------
                # NEXT QUESTION
                # -----------------------------------

                if st.button(
                    "Next Question ➜",
                    type="primary"
                ):

                    try:

                        with st.spinner(
                            "Generating adaptive question…"
                        ):

                            next_question = (
                                generate_verified_assessment_question(
                                    source_scope,
                                    question_type,
                                    st.session_state.quiz_difficulty,
                                    topic_scope
                                )
                            )

                            st.session_state.quiz_question = (
                                next_question
                            )

                            st.session_state.quiz_question_number += 1

                            st.session_state.quiz_answered = False

                        st.rerun()

                    except Exception as e:

                        st.error(
                            f"Next question could not "
                            f"be created: {describe_api_error(e)}"
                        )

        # -------------------------------------------
        # TOPIC PERFORMANCE
        # -------------------------------------------

        if (
            "topic_performance"
            in st.session_state
            and
            st.session_state.topic_performance
        ):

            st.divider()

            with st.expander(
                "📊 Topic Performance — full view in 📊 Mastery Dashboard"
            ):

                for topic, stats in (
                    st.session_state.topic_performance.items()
                ):

                    topic_accuracy = (
                        stats["correct"]
                        / stats["total"]
                    ) * 100

                    st.write(
                        f"**{topic}** — "
                        f"{stats['correct']}/"
                        f"{stats['total']} correct "
                        f"({topic_accuracy:.0f}%)"
                    )

                    st.progress(
                        topic_accuracy / 100
                    )

        st.divider()

        st.caption(
            "Reset Assessment clears only the current score and "
            "question; your learner history and Mastery are kept."
        )

        if st.button(
            "Reset Assessment"
        ):

            reset_quiz()
            save_learner_state()

            st.rerun()

        # ===================================================
# TAB 4 - MASTERY DASHBOARD
# ===================================================

with tab4:

    section_header(
        "📊 Mastery Dashboard",
        "Accuracy, topic Mastery, strengths, weak topics and likely "
        "misconceptions — computed only from your own answers.",
        eyebrow="Step 6 · Learner model"
    )

    learner_profile = st.session_state.get(
        "learner_profile",
        {}
    )

    topic_data = st.session_state.get(
        "topic_performance",
        {}
    )

    learner_mastery = st.session_state.get(
        "learner_mastery",
        {}
    )

    total_questions = st.session_state.get("quiz_total", 0)
    correct_answers = st.session_state.get("quiz_score", 0)

    if total_questions:
        overall_accuracy = (
            correct_answers / total_questions
        ) * 100
    else:
        overall_accuracy = 0

    # -------------------------------------------
    # TOPIC RESULTS (same rules as before:
    # Mastered = at least 2 attempts and >= 80 %,
    # Developing = >= 60 %, otherwise Needs Review)
    # -------------------------------------------

    topic_results = []

    for topic, stats in topic_data.items():

        topic_accuracy = (
            stats["correct"]
            / stats["total"]
        ) * 100

        if (
            stats["total"] >= 2
            and topic_accuracy >= 80
        ):

            status = "🟢 Mastered"

        elif topic_accuracy >= 60:

            status = "🟡 Developing"

        else:

            status = "🔴 Needs Review"

        topic_results.append({
            "topic": topic,
            "correct": stats["correct"],
            "total": stats["total"],
            "accuracy": topic_accuracy,
            "status": status
        })

    # Sort weakest topic first
    topic_results.sort(
        key=lambda item: item["accuracy"]
    )

    mastered_topics = sum(
        1 for result in topic_results
        if result["status"] == "🟢 Mastered"
    )

    developing_topics = sum(
        1 for result in topic_results
        if result["status"] == "🟡 Developing"
    )

    weak_topics = sum(
        1 for result in topic_results
        if result["status"] == "🔴 Needs Review"
    )

    weakest_topic = topic_results[0] if topic_results else None

    strongest_topic = (
        max(
            topic_results,
            key=lambda item: item["accuracy"]
        )
        if topic_results
        else None
    )

    tracked_topics = set(topic_data) | set(learner_mastery)

    # -------------------------------------------
    # SUMMARY CARDS
    # -------------------------------------------

    summary1, summary2, summary3, summary4 = st.columns(4)

    with summary1:

        st.metric(
            "Overall Accuracy",
            f"{overall_accuracy:.0f}%" if total_questions else "—"
        )

        st.caption(
            f"{correct_answers} / {total_questions} correct"
        )

    with summary2:

        st.metric(
            "Topics Tracked",
            len(tracked_topics)
        )

        st.caption(
            "Next difficulty: "
            + str(st.session_state.get("quiz_difficulty", "Medium"))
        )

    with summary3:

        st.metric(
            "Strongest Topic",
            strongest_topic["topic"] if strongest_topic else "—"
        )

        if strongest_topic:
            st.caption(
                f"{strongest_topic['accuracy']:.0f}% accuracy"
            )

    with summary4:

        st.metric(
            "Recommended Focus",
            weakest_topic["topic"] if weakest_topic else "—"
        )

        if weakest_topic:
            st.caption(
                f"{weakest_topic['accuracy']:.0f}% accuracy"
            )

    if learner_profile.get("diagnostic_completed", False):

        diagnostic_text = (
            f"{learner_profile.get('diagnostic_score', 0)}/"
            f"{learner_profile.get('diagnostic_total', 3)}"
        )

    else:

        diagnostic_text = "not completed"

    st.caption(
        "Learner level: "
        + str(learner_profile.get("level", "Intermediate"))
        + " · Explanation style: "
        + str(learner_profile.get("explanation_style", "Balanced"))
        + " · Diagnostic: "
        + diagnostic_text
    )

    if learner_profile.get("goal"):

        st.caption(
            "Learning goal: " + str(learner_profile["goal"])
        )

    st.divider()

    if not total_questions:

        st.info(
            "No assessment history yet. Answer a few questions in "
            "🧠 Adaptive Assessment to see accuracy, topic "
            "performance and a recommended focus here."
        )

    else:

        # -------------------------------------------
        # OVERALL ACCURACY
        # -------------------------------------------

        st.subheader("🎯 Overall Accuracy")

        st.progress(
            min(overall_accuracy / 100, 1.0)
        )

        if total_questions < 3:

            st.info(
                "More assessment answers are needed for a reliable "
                "Mastery estimate."
            )

        elif overall_accuracy >= 80:

            st.success(
                "Strong overall performance."
            )

        elif overall_accuracy >= 60:

            st.warning(
                "Good progress. Some topics still need reinforcement."
            )

        else:

            st.error(
                "More practice is recommended before increasing difficulty."
            )

        # -------------------------------------------
        # TOPIC PERFORMANCE
        # -------------------------------------------

        st.subheader("📚 Topic Performance")

        if not topic_results:

            st.info(
                "No topic performance data is available yet."
            )

        else:

            with st.container(horizontal=True, gap="small"):

                st.badge(
                    f"Mastered: {mastered_topics}",
                    icon=":material/verified:",
                    color="green"
                )

                st.badge(
                    f"Developing: {developing_topics}",
                    icon=":material/trending_up:",
                    color="orange"
                )

                st.badge(
                    f"Needs Review: {weak_topics}",
                    icon=":material/flag:",
                    color="red"
                )

            status_colors = {
                "🟢 Mastered": "green",
                "🟡 Developing": "orange",
                "🔴 Needs Review": "red"
            }

            for result_index, result in enumerate(topic_results):

                with st.container(key=f"card_topic_{result_index}"):

                    col_topic, col_status = st.columns(
                        [3, 1],
                        vertical_alignment="center"
                    )

                    with col_topic:

                        st.markdown(
                            f"**{result['topic']}**"
                        )

                        st.progress(
                            result["accuracy"] / 100,
                            text=(
                                f"{result['correct']} / "
                                f"{result['total']} correct"
                            )
                        )

                    with col_status:

                        st.badge(
                            result["status"],
                            color=status_colors[result["status"]]
                        )

                        st.caption(
                            f"{result['accuracy']:.0f}% accuracy"
                        )

            # ---------------------------------------
            # STUDY RECOMMENDATION
            # ---------------------------------------

            st.subheader(
                "💡 Recommended Next Step"
            )

            if weakest_topic["accuracy"] < 60:

                st.write(
                    f"Review **{weakest_topic['topic']}** "
                    "before continuing with harder questions. "
                    "Open the 🎓 Unified Tutor and ask "
                    "LearnLens for an explanation or examples "
                    "about this topic."
                )

            elif weakest_topic["accuracy"] < 80:

                st.write(
                    f"Continue practicing "
                    f"**{weakest_topic['topic']}**. "
                    "Your understanding is developing, but "
                    "additional questions would improve Mastery."
                )

            else:

                st.write(
                    "Your current topic performance is strong. "
                    "Continue with harder Adaptive Assessment questions "
                    "to test deeper understanding."
                )

    # -------------------------------------------
    # LEARNER MASTERY MODEL (EMA)
    # -------------------------------------------

    st.divider()

    st.subheader("🧩 Learner Mastery")

    st.caption(
        "Per-topic Mastery estimate (0.0–1.0) updated with an "
        "exponential moving average: new = old × (1 − α) + "
        "evidence × α. New topics start at 0.50. Assessment and "
        "diagnostic answers: evidence 1.0 / 0.0, α = 0.30. Tutor "
        "conversations: conversation-based heuristic evidence 0.55, "
        "α = 0.05 (0.35, α = 0.08 when the question signals difficulty)."
    )

    if not learner_mastery:

        st.info(
            "No Mastery data yet. Answer assessment or diagnostic "
            "questions, or ask the 🎓 Unified Tutor, to start "
            "building your learner model."
        )

    else:

        conversation_counts = st.session_state.get(
            "conversation_counts",
            {}
        )

        mastery_rows = []

        for topic, mastery in sorted(
            learner_mastery.items(),
            key=lambda item: item[1]
        ):

            quiz_stats = topic_data.get(
                topic,
                {}
            )

            mastery_rows.append({
                "Topic": topic,
                "Mastery": float(mastery),
                "Status": mastery_status_label(
                    float(mastery)
                ),
                "Quiz": (
                    f"{quiz_stats.get('correct', 0)}/"
                    f"{quiz_stats.get('total', 0)}"
                ),
                "Tutor Conversations": int(
                    conversation_counts.get(topic, 0)
                )
            })

        st.dataframe(
            mastery_rows,
            hide_index=True,
            column_config={
                "Mastery": st.column_config.ProgressColumn(
                    "Mastery",
                    min_value=0.0,
                    max_value=1.0,
                    format="%.2f"
                )
            }
        )

    # -------------------------------------------
    # LIKELY MISCONCEPTIONS
    # -------------------------------------------

    st.divider()

    st.subheader("🧠 Likely Misconceptions")

    st.caption(
        "Built only from your incorrect assessment answers "
        "(latest 3). Each card shows what you answered next to "
        "the source-grounded correct answer."
    )

    incorrect_history = [
        item
        for item in st.session_state.get(
            "quiz_history",
            []
        )
        if isinstance(item, dict)
        and not item.get("is_correct", True)
    ][-3:]

    if not incorrect_history:

        st.info(
            "No incorrect assessment answers yet — no "
            "misconceptions detected."
        )

    else:

        for card_index, item in enumerate(reversed(incorrect_history)):

            with st.container(key=f"card_misconception_{card_index}"):

                st.caption(
                    "Topic: " + str(item.get("topic", "Unknown"))
                )

                st.markdown(
                    "**" + str(item.get("question", "")) + "**"
                )

                answer_col, correct_col = st.columns(2)

                with answer_col:

                    st.caption("Your answer")

                    st.write(
                        str(item.get("student_answer", ""))
                    )

                with correct_col:

                    st.caption("Correct answer")

                    st.write(
                        str(item.get("correct_answer", ""))
                    )

                if item.get("explanation"):

                    st.caption(
                        "💡 " + str(item["explanation"])
                    )

                st.caption(
                    "📚 Source: " + str(item.get("source", "Unknown"))
                )

# ===================================================
# TAB 5 - VIDEO UNDERSTANDING
# ===================================================

with tab5:

    section_header(
        "🎞️ Video Learning Materials",
        "Upload a lecture recording. Gemini splits it into "
        "timestamped segments that the tutor can cite.",
        eyebrow="Step 2 · Learning materials"
    )

    with st.container(key="card_upload_video"):

        uploaded_video = st.file_uploader(
            "Upload an educational video",
            type=[
                "mp4",
                "mov",
                "avi",
                "mpeg",
                "mpg",
                "webm",
                "wmv"
            ],
            key="video_uploader"
        )

    if uploaded_video:

        video_col, _ = st.columns([3, 2])

        with video_col:

            st.video(
                uploaded_video.getvalue()
            )

        st.caption(
            f"File: {uploaded_video.name}"
        )

        # -------------------------------------------
        # ADD VIDEO TO KNOWLEDGE BASE
        # -------------------------------------------

        video_ref = None

        video_failure = st.session_state.get(
            "video_kb_failure"
        )

        if (
            video_failure
            and video_failure.get("name") == uploaded_video.name
        ):

            # Do not call Gemini again on every rerun after a
            # failure; retry only on request.
            st.error(
                "Video could not be added to the knowledge base: "
                + video_failure.get("message", "")
            )

            if st.button(
                "Retry video processing",
                key="retry_video_kb"
            ):
                st.session_state.pop("video_kb_failure", None)
                st.rerun()

        else:

            try:

                with st.spinner(
                    "Building semantic knowledge base from video…"
                ):

                    video_ref = upload_video_to_gemini(
                        uploaded_video
                    )

                    if (
                        "video_chunks" not in st.session_state
                        or st.session_state.get("video_kb_name")
                        != uploaded_video.name
                    ):

                        video_chunks = (
                            create_video_knowledge_chunks(
                                video_ref,
                                uploaded_video.name
                            )
                        )

                        if not video_chunks:
                            raise ValueError(
                                "Gemini returned no usable "
                                "video segments."
                            )

                        video_embeddings = (
                            create_embeddings(
                                video_chunks
                            )
                        )

                        st.session_state.video_chunks = (
                            video_chunks
                        )

                        st.session_state.video_embeddings = (
                            video_embeddings
                        )

                        st.session_state.video_kb_name = (
                            uploaded_video.name
                        )

                        st.rerun()

                if (
                    "video_chunks" in st.session_state
                    and st.session_state.get("video_kb_name")
                    == uploaded_video.name
                ):

                    status_badges([
                        "Video processed by Gemini",
                        f"{len(st.session_state.video_chunks)} "
                        "timestamped segments",
                        "Embeddings ready",
                        "Added to shared knowledge base"
                    ])

                    with st.expander("🎬 Timestamped segments"):

                        for video_segment in st.session_state.video_chunks:

                            st.markdown(
                                f"**{video_segment.get('start', '?')} – "
                                f"{video_segment.get('end', '?')}** · "
                                f"{video_segment.get('topic', '')}"
                            )

                            st.caption(
                                str(video_segment.get("text", ""))
                            )

            except Exception as e:

                for stale_key in (
                    "video_chunks",
                    "video_embeddings",
                    "video_kb_name"
                ):
                    st.session_state.pop(stale_key, None)

                st.session_state.video_kb_failure = {
                    "name": uploaded_video.name,
                    "message": describe_api_error(e)
                }

                st.rerun()

        st.divider()

        # -------------------------------------------
        # VIDEO QUESTION
        # -------------------------------------------

        st.subheader(
            "💬 Ask About This Video"
        )

        video_question = st.text_input(
            "Ask a question about the video:",
            key="video_question"
        )

        if st.button(
            "Ask Video",
            type="primary",
            key="ask_video"
        ):

            if not video_question.strip():

                st.warning(
                    "Please enter a question."
                )

            elif video_ref is None:

                st.warning(
                    "The video is not ready. Use "
                    "\"Retry video processing\" above first."
                )

            else:

                try:

                    with st.spinner(
                        "Analyzing video…"
                    ):

                        answer = (
                            answer_video_question(
                                video_ref,
                                video_question
                            )
                        )

                    with st.container(key="card_video_answer"):

                        eyebrow_label("Answer")

                        st.write(answer)

                except Exception as e:

                    st.error(
                        f"Video analysis failed: {describe_api_error(e)}"
                    )

    else:

        st.info(
            "No video uploaded yet. Short lecture clips work best; "
            "processing uses Gemini and can take a minute."
        )

        # Uploader emptied: remove the video part of the
        # knowledge base so stale material is not used.
        for stale_key in (
            "video_chunks",
            "video_embeddings",
            "video_kb_name",
            "video_kb_failure"
        ):
            st.session_state.pop(stale_key, None)

# ===================================================
# TAB 6 - SLIDE UNDERSTANDING
# ===================================================

with tab6:

    section_header(
        "📽️ Slide Learning Materials",
        "Upload a PowerPoint deck. LearnLens indexes slide text, "
        "tables and embedded images, and cites slide numbers.",
        eyebrow="Step 2 · Learning materials"
    )

    with st.container(key="card_upload_slides"):

        uploaded_pptx = st.file_uploader(
            "Upload PowerPoint presentation",
            type=["pptx"],
            key="pptx_uploader"
        )

    if uploaded_pptx:

        pptx_bytes = uploaded_pptx.getvalue()

        pptx_signature = hashlib.sha256(
            pptx_bytes
        ).hexdigest()

        slides_failure = st.session_state.get(
            "slides_index_failure"
        )

        if (
            slides_failure
            and slides_failure.get("signature") == pptx_signature
        ):

            # Do not call the embedding API again on every
            # rerun after a failure; retry only on request.
            st.error(
                "Presentation could not be processed: "
                + slides_failure.get("message", "")
            )

            if st.button(
                "Retry presentation processing",
                key="retry_slides_index"
            ):
                st.session_state.pop("slides_index_failure", None)
                st.rerun()

        elif (
            st.session_state.get(
                "pptx_signature"
            )
            != pptx_signature
        ):

            try:

                with st.spinner(
                    "Indexing learning materials…"
                ):

                    slide_chunks = (
                        extract_pptx_slides(
                            uploaded_pptx.name,
                            pptx_bytes
                        )
                    )

                    if not slide_chunks:
                        raise ValueError(
                            "The presentation contains no slides."
                        )

                    slide_embeddings = (
                        create_embeddings(
                            slide_chunks
                        )
                    )

                    st.session_state.pptx_signature = (
                        pptx_signature
                    )

                    st.session_state.slide_chunks = (
                        slide_chunks
                    )

                    st.session_state.slide_embeddings = (
                        slide_embeddings
                    )

                    st.rerun()

            except Exception as e:

                for stale_key in (
                    "pptx_signature",
                    "slide_chunks",
                    "slide_embeddings"
                ):
                    st.session_state.pop(stale_key, None)

                st.session_state.slides_index_failure = {
                    "signature": pptx_signature,
                    "message": describe_api_error(e)
                }

                st.rerun()

        slides_index_ready = (
            st.session_state.get("pptx_signature")
            == pptx_signature
            and "slide_chunks" in st.session_state
        )

        if slides_index_ready:

            slide_chunks = (
                st.session_state.slide_chunks
            )

            slide_embeddings = (
                st.session_state.slide_embeddings
            )

            image_count = sum(
                len(slide["images"])
                for slide in slide_chunks
            )

            status_badges([
                "1 presentation indexed",
                f"{len(slide_chunks)} slides",
                f"{image_count} embedded image(s)",
                "Embeddings ready"
            ])

            with st.expander(
                "📑 View extracted slide content"
            ):

                for slide in slide_chunks:

                    st.markdown(
                        f"**Slide {slide['slide']}**"
                    )

                    st.write(
                        slide["text"]
                    )

                    if slide["images"]:

                        st.caption(
                            f"{len(slide['images'])} "
                            "embedded image(s)"
                        )

                    st.divider()

            st.divider()

            st.subheader("💬 Ask About This Presentation")

            slide_question = st.text_input(
                "Ask a question about the presentation:",
                key="slide_question"
            )

            if st.button(
                "Ask Slides",
                type="primary"
            ):

                if not slide_question.strip():

                    st.warning(
                        "Please enter a question."
                    )

                else:

                    try:

                        with st.spinner(
                            "Searching learning materials…"
                        ):

                            (
                                slide_answer,
                                retrieved_slides
                            ) = answer_slide_question(
                                slide_question,
                                slide_chunks,
                                slide_embeddings
                            )

                        with st.container(key="card_slide_answer"):

                            eyebrow_label("Answer")

                            st.write(
                                slide_answer
                            )

                        eyebrow_label("Sources used")

                        st.caption(
                            " · ".join(
                                f"[{slide['file']}, Slide {slide['slide']}]"
                                for slide in retrieved_slides
                            )
                        )

                        with st.expander(
                            "📚 Retrieved slides"
                        ):

                            for slide in (
                                retrieved_slides
                            ):

                                st.markdown(
                                    f"**Slide "
                                    f"{slide['slide']}**"
                                )

                                st.write(
                                    slide["text"]
                                )

                                st.caption(
                                    f"Similarity score: "
                                    f"{slide['score']:.3f}"
                                )

                                st.divider()

                    except Exception as e:

                        st.error(
                            f"Slide analysis failed: {describe_api_error(e)}"
                        )

    else:

        st.info(
            "No slides uploaded yet. Upload a .pptx deck — for the "
            "benchmark demo use Kümeleme.pptx."
        )

        # Uploader emptied: remove the slide part of the
        # knowledge base so stale material is not used.
        for stale_key in (
            "pptx_signature",
            "slide_chunks",
            "slide_embeddings",
            "slides_index_failure"
        ):
            st.session_state.pop(stale_key, None)

# ===================================================
# TAB 7 - UNIFIED MULTIMODAL TUTOR
# ===================================================

with tab7:

    section_header(
        "🎓 Personalized Tutor",
        "Ask questions across all of your uploaded learning materials.",
        eyebrow="Step 4 · Grounded tutoring"
    )

    unified_items, unified_vectors = (
        build_unified_knowledge_base()
    )

    # -------------------------------------------
    # AVAILABLE SOURCES
    # -------------------------------------------

    available_sources = set()

    for item in unified_items:
        available_sources.add(
            item["source_type"]
        )

    source_labels = []

    if "pdf" in available_sources:
        source_labels.append("PDF")

    if "slide" in available_sources:
        source_labels.append("Slides")

    if "video" in available_sources:
        source_labels.append("Video")

    tutor_profile = st.session_state.get(
        "learner_profile",
        {}
    )

    with st.container(horizontal=True, gap="small"):

        st.badge(
            "Level: "
            + str(tutor_profile.get("level", "Intermediate")),
            icon=":material/school:",
            color="violet"
        )

        st.badge(
            "Style: "
            + str(tutor_profile.get("explanation_style", "Balanced")),
            icon=":material/tune:",
            color="blue"
        )

        st.badge(
            (
                f"{len(unified_items)} knowledge units · "
                + ", ".join(source_labels)
            )
            if unified_items
            else "No learning materials yet",
            icon=":material/library_books:",
            color="green" if unified_items else "gray"
        )

        st.badge(
            "Grounded mode",
            icon=":material/verified_user:",
            color="gray"
        )

    if not unified_items:

        st.info(
            "No learning materials yet. Upload a PDF (📄 PDF RAG), "
            "a presentation (📽️ Slide Understanding) or a video "
            "(🎞️ Video Understanding) first — everything joins "
            "this shared knowledge base."
        )

    else:

        # ---------------------------------------
        # ASK
        # ---------------------------------------

        with st.container(key="card_tutor_ask"):

            question_col, scope_col = st.columns(
                [3, 1],
                vertical_alignment="bottom"
            )

            with question_col:

                question = st.text_input(
                    "Ask LearnLens a question:",
                    placeholder="e.g. Minkowski uzaklığı nedir?",
                    key="unified_question"
                )

            with scope_col:

                search_scope = st.selectbox(
                    "Search scope",
                    [
                        "All Materials",
                        "PDF",
                        "Slides",
                        "Video"
                    ],
                    key="unified_scope"
                )

            ask_clicked = st.button(
                "Ask LearnLens",
                type="primary",
                key="ask_unified",
                icon=":material/send:"
            )

            st.caption(
                "Answers use only your uploaded learning materials and "
                "cite page / slide / timestamp. Questions outside the "
                "materials are declined instead of guessed."
            )

        if ask_clicked:

            if not question.strip():

                st.warning(
                    "Please enter a question."
                )

            else:

                filtered_items = []
                filtered_vectors = []

                for index, item in enumerate(
                    unified_items
                ):

                    include = False

                    if search_scope == "All Materials":
                        include = True

                    elif (
                        search_scope == "PDF"
                        and item["source_type"] == "pdf"
                    ):
                        include = True

                    elif (
                        search_scope == "Slides"
                        and item["source_type"] == "slide"
                    ):
                        include = True

                    elif (
                        search_scope == "Video"
                        and item["source_type"] == "video"
                    ):
                        include = True

                    if include:

                        filtered_items.append(
                            item
                        )

                        filtered_vectors.append(
                            unified_vectors[index]
                        )

                if not filtered_items:

                    st.warning(
                        "No learning materials are available "
                        "for the selected scope."
                    )

                else:

                    filtered_vectors = np.array(
                        filtered_vectors,
                        dtype=np.float32
                    )

                    try:

                        with st.spinner(
                            "Searching learning materials…"
                        ):

                            retrieved_sources = (
                                retrieve_unified_sources(
                                    question,
                                    filtered_items,
                                    filtered_vectors,
                                    min(
                                        6,
                                        len(filtered_items)
                                    )
                                )
                            )

                        with st.spinner(
                            "Generating grounded answer…"
                        ):

                            answer = (
                                answer_unified_question(
                                    question,
                                    retrieved_sources,
                                    st.session_state.get(
                                        "learner_profile",
                                        {}
                                    ),
                                    st.session_state.get(
                                        "topic_performance",
                                        {}
                                    )
                                )
                            )

                        answer_refused = is_refusal_answer(answer)

                        if answer_refused:

                            # Hallucination control is a product
                            # feature, shown calmly (not an error).
                            with st.container(key="card_tutor_refusal"):

                                st.markdown(
                                    "**🔍 Not found in your learning "
                                    "materials**"
                                )

                                st.write(answer)

                                st.caption(
                                    "LearnLens answers only from your "
                                    "uploaded learning materials, so it "
                                    "declines instead of guessing. The "
                                    "learner model is not updated."
                                )

                        else:

                            with st.container(key="card_tutor_answer"):

                                eyebrow_label("Answer")

                                st.write(answer)

                        # ---------------------------
                        # KNOWLEDGE TAGS FOR THE USED
                        # SOURCES (only missing ones,
                        # cached; never blocks answers)
                        # ---------------------------

                        if not answer_refused:

                            with st.spinner(
                                "Updating knowledge structure…"
                            ):

                                (
                                    new_tag_count,
                                    failed_tag_count
                                ) = auto_tag_knowledge_items(
                                    retrieved_sources
                                )

                            if new_tag_count:

                                st.caption(
                                    f"🧩 Tagged {new_tag_count} new "
                                    "knowledge unit(s) (cached in "
                                    "knowledge_tags.json)."
                                )

                            elif failed_tag_count:

                                st.caption(
                                    "🧩 Knowledge tags could not be "
                                    "generated for these sources; the "
                                    "file name is used as topic."
                                )

                        # ---------------------------
                        # LEARNER MODEL UPDATE
                        # (conversation-based heuristic
                        # evidence, small EMA step)
                        # ---------------------------

                        try:

                            conversation_update = (
                                record_conversation_mastery(
                                    question,
                                    retrieved_sources,
                                    answer
                                )
                            )

                            if conversation_update:

                                save_learner_state()

                                st.caption(
                                    "🧩 Learner model updated "
                                    "(conversation heuristic): "
                                    f"{conversation_update['topic']} → "
                                    f"{conversation_update['mastery']:.2f}"
                                    + (
                                        " · difficulty signal detected"
                                        if conversation_update[
                                            "struggle"
                                        ]
                                        else ""
                                    )
                                )

                        except Exception:

                            # The answer is already shown;
                            # a learner-model hiccup must not
                            # break tutoring.
                            pass

                        # ---------------------------
                        # SOURCES
                        # ---------------------------

                        eyebrow_label("Sources used")

                        st.caption(
                            " · ".join(
                                source["citation"]
                                for source in retrieved_sources
                            )
                        )

                        with st.expander(
                            f"📚 Retrieved passages ({len(retrieved_sources)})"
                        ):

                            for source in retrieved_sources:

                                st.markdown(
                                    f"**{source['citation']}**"
                                )

                                st.caption(
                                    f"Source type: "
                                    f"{source['source_type']}"
                                    + (
                                        f" · Topic: {source_topic_tag(source)}"
                                        if source_topic_tag(source)
                                        else ""
                                    )
                                    + f" · Similarity: {source['score']:.3f}"
                                )

                                st.write(
                                    source["text"]
                                )

                                st.divider()

                    except Exception as e:

                        st.error(
                            f"Unified tutor failed: {describe_api_error(e)}"
                        )

        # ---------------------------------------
        # KNOWLEDGE STRUCTURE
        # ---------------------------------------

        st.divider()

        with st.expander("🧩 Knowledge Structure"):

            tagged_units = sum(
                1 for item in unified_items
                if item.get("tagged")
            )

            untagged_units = (
                len(unified_items) - tagged_units
            )

            structure_items = [
                item for item in unified_items
                if item.get("tagged")
                or str(item.get("topic", "")).strip()
            ]

            structure_topics = {}

            for item in structure_items:

                structure_topic = (
                    source_topic_tag(item) or "Untitled topic"
                )

                structure_topics.setdefault(
                    structure_topic,
                    []
                ).append(item)

            ks_col1, ks_col2 = st.columns(2)

            with ks_col1:

                st.metric(
                    "Tagged Units",
                    f"{tagged_units} / {len(unified_items)}"
                )

            with ks_col2:

                st.metric(
                    "Topics",
                    len(structure_topics)
                )

            st.caption(
                "Each unit gets a topic, subtopic, key concepts and "
                "prerequisites. Tags are generated in batches of "
                f"{KNOWLEDGE_TAG_BATCH_SIZE} units per Gemini call and "
                "cached in knowledge_tags.json — cached units are never "
                "re-tagged."
            )

            st.toggle(
                "Auto-tag missing units when they are used",
                value=True,
                key="auto_tagging_enabled",
                help=(
                    "When on, the tutor and the assessment tag only "
                    "the few untagged units they actually use (at "
                    "most one batched Gemini call). Turn off to save "
                    "API quota; tagging then happens only via the "
                    "button below."
                )
            )

            if untagged_units:

                if st.button(
                    f"Generate tags for all {untagged_units} "
                    "untagged unit(s)",
                    key="generate_knowledge_tags"
                ):

                    tag_progress = st.progress(
                        0.0,
                        text="Building knowledge structure…"
                    )

                    def update_tag_progress(done, total):
                        tag_progress.progress(
                            done / total if total else 1.0,
                            text=(
                                f"Tagged batch: {done} / "
                                f"{total} units"
                            )
                        )

                    (
                        tagged_count,
                        failed_count,
                        pending_count
                    ) = tag_knowledge_items(
                        unified_items,
                        update_tag_progress
                    )

                    if failed_count:

                        st.session_state.knowledge_tag_message = (
                            f"Tagged {tagged_count} unit(s); "
                            f"{failed_count} could not be tagged "
                            "(the tutor still works without tags)."
                        )

                    else:

                        st.session_state.knowledge_tag_message = (
                            f"Tagged {tagged_count} unit(s)."
                        )

                    st.rerun()

            if st.session_state.get("knowledge_tag_message"):

                st.info(
                    st.session_state.pop(
                        "knowledge_tag_message"
                    )
                )

            if not structure_topics:

                st.info(
                    "Generate knowledge structure to organize your "
                    "uploaded learning materials into topics, concepts "
                    "and prerequisites."
                )

            else:

                selected_structure_topic = st.selectbox(
                    "Browse a topic",
                    sorted(structure_topics),
                    format_func=lambda name: (
                        f"{name} ({len(structure_topics[name])} units)"
                    ),
                    key="knowledge_topic_filter"
                )

                for unit_index, unit in enumerate(
                    structure_topics.get(
                        selected_structure_topic,
                        []
                    )
                ):

                    with st.container(key=f"card_unit_{unit_index}"):

                        st.markdown(
                            f"**{unit['citation']}**"
                        )

                        if unit.get("subtopic"):

                            st.caption(
                                "Subtopic: " + str(unit["subtopic"])
                            )

                        st.write(
                            "Key concepts: "
                            + (
                                ", ".join(unit.get("key_concepts", []))
                                or "—"
                            )
                        )

                        st.write(
                            "Prerequisites: "
                            + (
                                ", ".join(unit.get("prerequisites", []))
                                or "—"
                            )
                        )

# ===================================================
# TAB 8 - LEARNER PROFILE / COLD START
# ===================================================

with tab8:

    section_header(
        "🧭 Learner Profile",
        "Tell LearnLens about your level and learning goal. "
        "Your profile personalizes tutoring and Adaptive Assessment.",
        eyebrow="Step 1 · Personalization"
    )

    if "learner_profile" not in st.session_state:

        st.session_state.learner_profile = {
            "level": "Intermediate",
            "goal": "",
            "explanation_style": "Balanced",
            "diagnostic_completed": False,
            "diagnostic_score": 0,
            "diagnostic_total": 0
        }

    profile = st.session_state.learner_profile

    # Fill missing / invalid values (old or hand-edited
    # learner_state.json) so the widgets below cannot fail.
    for default_key, default_value in {
        "level": "Intermediate",
        "goal": "",
        "explanation_style": "Balanced",
        "diagnostic_completed": False,
        "diagnostic_score": 0,
        "diagnostic_total": 0
    }.items():
        profile.setdefault(default_key, default_value)

    if profile["level"] not in (
        "Beginner",
        "Intermediate",
        "Advanced"
    ):
        profile["level"] = "Intermediate"

    if profile["explanation_style"] not in (
        "Concise",
        "Balanced",
        "Detailed"
    ):
        profile["explanation_style"] = "Balanced"

    if not isinstance(profile["goal"], str):
        profile["goal"] = ""

    # After the diagnostic changes the level, rebuild the
    # profile widgets from the saved profile.
    if st.session_state.pop("sync_profile_widgets", False):

        for widget_key in (
            "profile_level",
            "profile_goal",
            "profile_explanation_style"
        ):
            st.session_state.pop(widget_key, None)

    profile_col, summary_col = st.columns([3, 2], gap="large")

    with profile_col:

        with st.container(key="card_profile_form"):

            eyebrow_label("Your learning profile")

            learner_level = st.selectbox(
                "How would you describe your current level?",
                [
                    "Beginner",
                    "Intermediate",
                    "Advanced"
                ],
                index=[
                    "Beginner",
                    "Intermediate",
                    "Advanced"
                ].index(
                    profile.get(
                        "level",
                        "Intermediate"
                    )
                ),
                key="profile_level"
            )

            learner_goal = st.text_area(
                "What is your learning goal?",
                value=profile.get(
                    "goal",
                    ""
                ),
                placeholder=(
                    "Example: I want to understand clustering "
                    "methods and solve exam questions."
                ),
                key="profile_goal"
            )

            explanation_style = st.selectbox(
                "Preferred explanation style",
                [
                    "Concise",
                    "Balanced",
                    "Detailed"
                ],
                index=[
                    "Concise",
                    "Balanced",
                    "Detailed"
                ].index(
                    profile.get(
                        "explanation_style",
                        "Balanced"
                    )
                ),
                key="profile_explanation_style"
            )

            if st.button(
                "Save Learning Profile",
                type="primary",
                key="save_learning_profile"
            ):

                st.session_state.learner_profile[
                    "level"
                ] = learner_level

                st.session_state.learner_profile[
                    "goal"
                ] = learner_goal.strip()

                st.session_state.learner_profile[
                    "explanation_style"
                ] = explanation_style

                save_learner_state()

                st.success(
                    "✅ Learning profile saved."
                )

    current_profile = (
        st.session_state.learner_profile
    )

    with summary_col:

        with st.container(key="card_profile_summary"):

            eyebrow_label("Current profile")

            col1, col2 = st.columns(2)

            with col1:

                st.metric(
                    "Current Level",
                    current_profile[
                        "level"
                    ]
                )

            with col2:

                st.metric(
                    "Explanation Style",
                    current_profile[
                        "explanation_style"
                    ]
                )

            diagnostic_status = (
                "Completed"
                if current_profile[
                    "diagnostic_completed"
                ]
                else "Not completed"
            )

            st.metric(
                "Diagnostic",
                diagnostic_status
            )

            if current_profile["goal"]:

                st.caption("Learning goal")

                st.write(
                    current_profile["goal"]
                )

            else:

                st.caption(
                    "No learning goal has been entered yet."
                )

    st.divider()

    section_header(
        "🧪 Starting Level Diagnostic",
        "Let LearnLens estimate your starting level using 3 "
        "questions from your own learning materials.",
        eyebrow="Step 3 · Cold start"
    )

    if not st.session_state.get("diagnostic_active", False):

        step1, step2, step3 = st.columns(3)

        for step_column, step_number, step_level in (
            (step1, 1, "Easy"),
            (step2, 2, "Medium"),
            (step3, 3, "Hard")
        ):

            with step_column:

                with st.container(key=f"card_diagnostic_step_{step_number}"):

                    st.caption(f"Question {step_number}")

                    difficulty_badge(step_level)

        st.caption(
            "Score 0–1 → Beginner / Easy · 2 → Intermediate / Medium · "
            "3 → Advanced / Hard. Each answer also updates your "
            "topic Mastery."
        )

    # -------------------------------------------
    # INITIALIZE DIAGNOSTIC STATE
    # -------------------------------------------

    if "diagnostic_active" not in st.session_state:
        st.session_state.diagnostic_active = False

    if "diagnostic_index" not in st.session_state:
        st.session_state.diagnostic_index = 0

    if "diagnostic_score" not in st.session_state:
        st.session_state.diagnostic_score = 0

    if "diagnostic_question" not in st.session_state:
        st.session_state.diagnostic_question = None

    if "diagnostic_answered" not in st.session_state:
        st.session_state.diagnostic_answered = False

    if "diagnostic_last_correct" not in st.session_state:
        st.session_state.diagnostic_last_correct = False

    diagnostic_difficulties = [
        "Easy",
        "Medium",
        "Hard"
    ]

    # -------------------------------------------
    # START DIAGNOSTIC
    # -------------------------------------------

    if not st.session_state.diagnostic_active:

        diagnostic_items, _ = build_unified_knowledge_base()

        if not diagnostic_items:

            st.warning(
                "Upload learning material first — diagnostic "
                "questions are generated from your own material."
            )

        if st.button(
            "Start Diagnostic",
            type="primary",
            key="start_diagnostic",
            disabled=not diagnostic_items
        ):

            st.session_state.diagnostic_active = True
            st.session_state.diagnostic_index = 0
            st.session_state.diagnostic_score = 0
            st.session_state.diagnostic_answered = False
            st.session_state.diagnostic_result_applied = False

            diagnostic_topic = (
                st.session_state.learner_profile.get(
                    "goal",
                    ""
                )
            )

            try:

                with st.spinner(
                    "Generating adaptive question…"
                ):

                    st.session_state.diagnostic_question = (
                        generate_verified_assessment_question(
                            "All Materials",
                            "MCQ",
                            "Easy",
                            diagnostic_topic
                        )
                    )

                st.rerun()

            except Exception as e:

                st.session_state.diagnostic_active = False

                st.error(
                    f"Diagnostic could not be started: {describe_api_error(e)}"
                )

    # -------------------------------------------
    # ACTIVE DIAGNOSTIC
    # -------------------------------------------

    else:

        diagnostic_number = (
            st.session_state.diagnostic_index + 1
        )

        st.progress(
            diagnostic_number / 3,
            text=(
                f"Diagnostic question {diagnostic_number} of 3"
            )
        )

        question_data = (
            st.session_state.diagnostic_question
        )

        if question_data is not None:

            with st.container(key="card_question_diagnostic"):

                with st.container(horizontal=True, gap="small"):

                    st.badge(
                        f"Question {diagnostic_number} / 3",
                        color="gray"
                    )

                    difficulty_badge(
                        diagnostic_difficulties[
                            st.session_state.diagnostic_index
                        ]
                    )

                st.markdown(
                    f"#### {question_data['question']}"
                )

                diagnostic_answer = st.radio(
                    "Choose your answer:",
                    question_data["options"],
                    index=None,
                    key=(
                        "diagnostic_answer_"
                        + str(
                            st.session_state.diagnostic_index
                        )
                    ),
                    disabled=(
                        st.session_state.diagnostic_answered
                    )
                )

            # -----------------------------------
            # SUBMIT
            # -----------------------------------

            if not st.session_state.diagnostic_answered:

                if st.button(
                    "Submit Diagnostic Answer",
                    type="primary",
                    key=(
                        "submit_diagnostic_"
                        + str(
                            st.session_state.diagnostic_index
                        )
                    )
                ):

                    if diagnostic_answer is None:

                        st.warning(
                            "Please select an answer."
                        )

                    else:

                        is_correct = (
                            diagnostic_answer
                            == question_data[
                                "correct_answer"
                            ]
                        )

                        if is_correct:

                            st.session_state.diagnostic_score += 1

                        # Diagnostic answers update the EMA
                        # learner model, but not topic_performance
                        # (kept separate from assessment accuracy).
                        record_quiz_mastery(
                            question_data.get(
                                "topic",
                                "Diagnostic"
                            ),
                            is_correct
                        )

                        save_learner_state()

                        st.session_state.diagnostic_last_correct = (
                            is_correct
                        )

                        st.session_state.diagnostic_answered = True

                        st.rerun()

            # -----------------------------------
            # RESULT
            # -----------------------------------

            else:

                if st.session_state.diagnostic_last_correct:

                    st.success(
                        "✓ Correct",
                        icon=":material/check_circle:"
                    )

                else:

                    st.error(
                        "✗ Needs Review",
                        icon=":material/error:"
                    )

                    st.write(
                        "**Correct answer:** "
                        + str(
                            question_data[
                                "correct_answer"
                            ]
                        )
                    )

                st.info(
                    "💡 "
                    + question_data[
                        "explanation"
                    ]
                )

                st.caption(
                    "📚 Source: "
                    + question_data[
                        "source_citation"
                    ]
                )

                # -------------------------------
                # QUESTIONS 1 AND 2
                # -------------------------------

                if (
                    st.session_state.diagnostic_index
                    < 2
                ):

                    if st.button(
                        "Next Diagnostic Question ➜",
                        type="primary",
                        key=(
                            "next_diagnostic_"
                            + str(
                                st.session_state.diagnostic_index
                            )
                        )
                    ):

                        next_index = (
                            st.session_state.diagnostic_index
                            + 1
                        )

                        next_difficulty = (
                            diagnostic_difficulties[
                                next_index
                            ]
                        )

                        diagnostic_topic = (
                            st.session_state.learner_profile.get(
                                "goal",
                                ""
                            )
                        )

                        try:

                            with st.spinner(
                                "Generating adaptive question…"
                            ):

                                next_question = (
                                    generate_verified_assessment_question(
                                        "All Materials",
                                        "MCQ",
                                        next_difficulty,
                                        diagnostic_topic
                                    )
                                )

                            st.session_state.diagnostic_index = (
                                next_index
                            )

                            st.session_state.diagnostic_question = (
                                next_question
                            )

                            st.session_state.diagnostic_answered = (
                                False
                            )

                            st.rerun()

                        except Exception as e:

                            st.error(
                                f"Next diagnostic question "
                                f"could not be created: {describe_api_error(e)}"
                            )

                # -------------------------------
                # QUESTION 3 FINISHED
                # -------------------------------

                else:

                    final_score = (
                        st.session_state.diagnostic_score
                    )

                    if final_score <= 1:

                        estimated_level = (
                            "Beginner"
                        )

                        starting_difficulty = (
                            "Easy"
                        )

                    elif final_score == 2:

                        estimated_level = (
                            "Intermediate"
                        )

                        starting_difficulty = (
                            "Medium"
                        )

                    else:

                        estimated_level = (
                            "Advanced"
                        )

                        starting_difficulty = (
                            "Hard"
                        )

                    # Apply the result ONCE. This block is
                    # re-rendered on every rerun; re-applying
                    # would reset the adaptive difficulty that
                    # the assessment changes afterwards.
                    if not st.session_state.get(
                        "diagnostic_result_applied",
                        False
                    ):

                        st.session_state.learner_profile[
                            "diagnostic_completed"
                        ] = True

                        st.session_state.learner_profile[
                            "diagnostic_score"
                        ] = final_score

                        st.session_state.learner_profile[
                            "diagnostic_total"
                        ] = 3

                        st.session_state.learner_profile[
                            "diagnostic_level"
                        ] = estimated_level

                        st.session_state.learner_profile[
                            "level"
                        ] = estimated_level

                        st.session_state.quiz_difficulty = (
                            starting_difficulty
                        )

                        st.session_state.quiz_correct_streak = 0

                        save_learner_state()

                        st.session_state.diagnostic_result_applied = True

                        # Refresh the profile widgets and the
                        # assessment metrics with the new level.
                        st.session_state.sync_profile_widgets = True

                        st.rerun()

                    st.success(
                        "🎯 Diagnostic completed — your starting "
                        "level is set."
                    )

                    result1, result2, result3 = (
                        st.columns(3)
                    )

                    with result1:

                        st.metric(
                            "Estimated Level",
                            estimated_level
                        )

                    with result2:

                        st.metric(
                            "Diagnostic Score",
                            f"{final_score}/3"
                        )

                    with result3:

                        st.metric(
                            "Starting Difficulty",
                            starting_difficulty
                        )

                    st.info(
                        "Next: open 🧠 Adaptive Assessment — it "
                        "starts at this difficulty and adapts as "
                        "you answer."
                    )

                    if st.button(
                        "Restart Diagnostic",
                        key="restart_diagnostic"
                    ):

                        st.session_state.diagnostic_active = False
                        st.session_state.diagnostic_index = 0
                        st.session_state.diagnostic_score = 0
                        st.session_state.diagnostic_question = None
                        st.session_state.diagnostic_answered = False
                        st.session_state.diagnostic_last_correct = False
                        st.session_state.diagnostic_result_applied = False

                        st.session_state.learner_profile[
                            "diagnostic_completed"
                        ] = False

                        save_learner_state()

                        st.rerun()

# ===================================================
# TAB 9 - SYSTEM EVALUATION
# ===================================================

with tab9:

    section_header(
        "🧪 Evaluation",
        "Evaluation measures whether LearnLens retrieves the right "
        "evidence, answers faithfully and refuses unsupported questions.",
        eyebrow="Step 7 · Evidence"
    )

    with st.container(key="card_metric_guide"):

        eyebrow_label("What the RAGAS metrics mean (0 – 1, higher is better)")

        guide1, guide2, guide3, guide4 = st.columns(4)

        for guide_column, guide_title, guide_text in (
            (
                guide1,
                "Faithfulness",
                "Is the answer supported by the retrieved evidence?"
            ),
            (
                guide2,
                "Answer Relevancy",
                "Does the answer address the question?"
            ),
            (
                guide3,
                "Context Precision",
                "Did retrieval return useful evidence?"
            ),
            (
                guide4,
                "Context Recall",
                "Did retrieval find the necessary evidence?"
            )
        ):

            with guide_column:

                st.markdown(f"**{guide_title}**")

                st.caption(guide_text)

    st.markdown("")

    section_header(
        "Real system metrics",
        "Measured on your uploaded learning materials with live "
        "Gemini calls.",
        eyebrow="Measured"
    )

    evaluation_items, evaluation_vectors = (
        build_unified_knowledge_base()
    )

    if not evaluation_items:

        st.info(
            "No learning materials yet. Upload materials first — "
            "for the benchmark test set, upload Kümeleme.pptx in "
            "📽️ Slide Understanding."
        )

    else:

        available_eval_sources = {
            item["source_type"]
            for item in evaluation_items
        }

        evaluation_source_options = [
            "All Materials"
        ]

        if "pdf" in available_eval_sources:
            evaluation_source_options.append(
                "PDF"
            )

        if "slide" in available_eval_sources:
            evaluation_source_options.append(
                "Slides"
            )

        if "video" in available_eval_sources:
            evaluation_source_options.append(
                "Video"
            )

        evaluation_scope = st.selectbox(
            "Evaluation source",
            evaluation_source_options,
            key="evaluation_scope"
        )

        evaluation_question = st.text_input(
            "Test question",
            placeholder=(
                "Example: Minkowski uzaklığı nedir?"
            ),
            key="evaluation_question"
        )

        reference_answer = st.text_area(
            "Reference answer",
            placeholder=(
                "Write the expected correct answer based "
                "on the uploaded material."
            ),
            key="evaluation_reference"
        )

        if st.button(
            "Run Evaluation",
            type="primary",
            key="run_evaluation"
        ):

            if not evaluation_question.strip():

                st.warning(
                    "Please enter a test question."
                )

            elif not reference_answer.strip():

                st.warning(
                    "Please enter a reference answer."
                )

            else:

                filtered_items = []
                filtered_vectors = []

                for index, item in enumerate(
                    evaluation_items
                ):

                    include = False

                    if (
                        evaluation_scope
                        == "All Materials"
                    ):
                        include = True

                    elif (
                        evaluation_scope == "PDF"
                        and item["source_type"] == "pdf"
                    ):
                        include = True

                    elif (
                        evaluation_scope == "Slides"
                        and item["source_type"] == "slide"
                    ):
                        include = True

                    elif (
                        evaluation_scope == "Video"
                        and item["source_type"] == "video"
                    ):
                        include = True

                    if include:

                        filtered_items.append(
                            item
                        )

                        filtered_vectors.append(
                            evaluation_vectors[index]
                        )

                if not filtered_items:

                    st.warning(
                        "No material is available for "
                        "the selected source."
                    )

                else:

                    filtered_vectors = np.array(
                        filtered_vectors,
                        dtype=np.float32
                    )

                    try:

                        with st.spinner(
                            "Searching learning materials…"
                        ):

                            retrieved_sources = (
                                retrieve_unified_sources(
                                    evaluation_question,
                                    filtered_items,
                                    filtered_vectors,
                                    min(
                                        6,
                                        len(filtered_items)
                                    )
                                )
                            )

                        with st.spinner(
                            "Generating grounded answer…"
                        ):

                            generated_answer = (
                                answer_unified_question(
                                    evaluation_question,
                                    retrieved_sources,
                                    st.session_state.get(
                                        "learner_profile",
                                        {}
                                    ),
                                    st.session_state.get(
                                        "topic_performance",
                                        {}
                                    )
                                )
                            )

                        retrieved_contexts = [
                            source["text"]
                            for source in retrieved_sources
                        ]

                        with st.spinner(
                            "Evaluating RAG quality…"
                        ):

                            scores = (
                                run_ragas_evaluation(
                                    evaluation_question,
                                    generated_answer,
                                    retrieved_contexts,
                                    reference_answer
                                )
                            )

                        st.subheader(
                            "📊 Evaluation Results"
                        )

                        col1, col2, col3, col4 = (
                            st.columns(4)
                        )

                        with col1:

                            st.metric(
                                "Faithfulness",
                                format_metric(clean_metric_value(scores.get('faithfulness')))
                            )

                        with col2:

                            st.metric(
                                "Answer Relevancy",
                                format_metric(clean_metric_value(scores.get('answer_relevancy')))
                            )

                        with col3:

                            st.metric(
                                "Context Precision",
                                format_metric(clean_metric_value(scores.get('context_precision')))
                            )

                        with col4:

                            st.metric(
                                "Context Recall",
                                format_metric(clean_metric_value(scores.get('context_recall')))
                            )

                        st.divider()

                        st.subheader(
                            "🎓 Generated Answer"
                        )

                        st.write(
                            generated_answer
                        )

                        with st.expander(
                            "🔎 Retrieved Evaluation Contexts"
                        ):

                            for source in retrieved_sources:

                                st.markdown(
                                    f"### {source['citation']}"
                                )

                                st.write(
                                    source["text"]
                                )

                                st.divider()

                    except Exception as e:

                        st.error(
                            f"Evaluation failed: {describe_api_error(e)}"
                        )

        # -------------------------------------------
        # BENCHMARK TEST SET (BATCH RAGAS)
        # -------------------------------------------

        st.divider()

        st.subheader("📚 Benchmark Test Set")

        st.write(
            "Run the team-built test set end-to-end: retrieval, "
            "grounded answer generation, RAGAS scoring for "
            "in-material questions and refusal checking for "
            "off-material questions."
        )

        try:

            benchmark_cases = load_evaluation_testset()
            testset_error = None

        except Exception as e:

            benchmark_cases = []
            testset_error = str(e)

        if testset_error:

            st.error(
                "Benchmark test set could not be loaded: "
                + testset_error
            )

        else:

            grounded_case_count = sum(
                1 for case in benchmark_cases
                if not case["should_refuse"]
            )

            st.caption(
                f"{len(benchmark_cases)} test cases "
                f"({grounded_case_count} in-material, "
                f"{len(benchmark_cases) - grounded_case_count} "
                "off-material) loaded from evaluation_testset.json"
            )

            with st.expander("View test cases"):

                st.dataframe(
                    [
                        {
                            "ID": case["id"],
                            "Question": case["question"],
                            "Expected Source": (
                                case["expected_source"] or "—"
                            ),
                            "Should Refuse": case["should_refuse"]
                        }
                        for case in benchmark_cases
                    ],
                    hide_index=True
                )

            # Warn when the material the test set expects is not
            # loaded: source matching would fail for every case.
            loaded_file_keys = {
                normalize_topic_key(item.get("file", ""))
                for item in evaluation_items
            }

            missing_expected_files = sorted({
                case["expected_source"].split(",")[0].strip()
                for case in benchmark_cases
                if case["expected_source"]
                and normalize_topic_key(
                    case["expected_source"].split(",")[0]
                ) not in loaded_file_keys
            })

            if missing_expected_files:

                st.warning(
                    "Expected source file(s) not in the knowledge "
                    "base: " + ", ".join(missing_expected_files)
                    + ". Upload them (same file name) or the "
                    "Expected Source Match will fail for those cases."
                )

            benchmark_scope = st.selectbox(
                "Benchmark source",
                evaluation_source_options,
                key="benchmark_scope"
            )

            benchmark_mode = st.radio(
                "Benchmark mode",
                [
                    "Quick Benchmark",
                    "Full Benchmark"
                ],
                index=0,
                horizontal=True,
                key="benchmark_mode",
                help=(
                    "Quick: 1 off-material case + first 2 "
                    "in-material cases (quota-safe). Full: every "
                    "case in evaluation_testset.json. RAGAS makes "
                    "several Gemini calls per case."
                )
            )

            if benchmark_mode == "Quick Benchmark":
                selected_benchmark_cases = (
                    select_quick_benchmark_cases(
                        benchmark_cases
                    )
                )
            else:
                selected_benchmark_cases = benchmark_cases

            st.caption(
                f"{len(selected_benchmark_cases)} case(s) will run, "
                f"with a {BENCHMARK_CASE_DELAY_SECONDS}s pause "
                "between cases."
            )

            if st.button(
                "Run Benchmark",
                type="primary",
                key="run_benchmark"
            ):

                benchmark_items, benchmark_vectors = (
                    filter_items_by_scope(
                        evaluation_items,
                        evaluation_vectors,
                        benchmark_scope
                    )
                )

                if not benchmark_items:

                    st.warning(
                        "No material is available for "
                        "the selected source."
                    )

                else:

                    benchmark_progress = st.progress(
                        0.0,
                        text="Starting benchmark..."
                    )

                    case_status_box = st.container()

                    benchmark_results = []

                    total_cases = len(selected_benchmark_cases)

                    quota_stopped = False

                    for number, case in enumerate(
                        selected_benchmark_cases,
                        start=1
                    ):

                        # Quota reached: record the rest as
                        # skipped, make no further API calls.
                        if quota_stopped:
                            benchmark_results.append(
                                skipped_benchmark_case(case)
                            )
                            continue

                        # Pause between cases to stay within
                        # free-tier rate limits.
                        if number > 1:
                            time.sleep(
                                BENCHMARK_CASE_DELAY_SECONDS
                            )

                        benchmark_progress.progress(
                            (number - 1) / total_cases,
                            text=(
                                f"Case {number}/{total_cases}: "
                                f"{case['question'][:70]}"
                            )
                        )

                        case_result = run_benchmark_case(
                            case,
                            benchmark_items,
                            benchmark_vectors,
                            st.session_state.get(
                                "learner_profile",
                                {}
                            ),
                            st.session_state.get(
                                "topic_performance",
                                {}
                            )
                        )

                        benchmark_results.append(case_result)

                        with case_status_box:

                            if case_result["status"] == "quota_error":

                                quota_stopped = True

                                st.write(
                                    f"⛔ Case {number} "
                                    f"({case['id']}) — quota_error"
                                )

                                st.error(QUOTA_ERROR_MESSAGE)

                            elif case_result["status"] != "ok":

                                st.write(
                                    f"⚠️ Case {number} "
                                    f"({case['id']}) — error: "
                                    f"{case_result['error']}"
                                )

                            elif case["should_refuse"]:

                                st.write(
                                    f"{'✅' if case_result['refusal_correct'] else '❌'} "
                                    f"Case {number} ({case['id']}) — "
                                    "refusal "
                                    + (
                                        "correct"
                                        if case_result["refusal_correct"]
                                        else "incorrect"
                                    )
                                )

                            else:

                                st.write(
                                    f"✅ Case {number} ({case['id']}) — "
                                    "F "
                                    f"{format_metric(case_result['faithfulness'])}"
                                    " · AR "
                                    f"{format_metric(case_result['answer_relevancy'])}"
                                    " · CP "
                                    f"{format_metric(case_result['context_precision'])}"
                                    " · CR "
                                    f"{format_metric(case_result['context_recall'])}"
                                )

                    benchmark_summary = summarize_benchmark(
                        benchmark_results
                    )

                    if benchmark_summary["successful_cases"] == 0:

                        if benchmark_summary["quota_exhausted"]:
                            final_progress_text = (
                                "Benchmark stopped — Gemini API "
                                "quota reached."
                            )
                        else:
                            final_progress_text = (
                                "Benchmark stopped — no case "
                                "completed successfully."
                            )

                    elif benchmark_summary["quota_exhausted"]:

                        final_progress_text = (
                            "Benchmark stopped early — Gemini API "
                            "quota reached (partial results)."
                        )

                    else:

                        final_progress_text = "Benchmark completed."

                    benchmark_progress.progress(
                        1.0,
                        text=final_progress_text
                    )

                    try:

                        st.session_state.benchmark_report = (
                            save_benchmark_results(
                                benchmark_results,
                                benchmark_summary,
                                benchmark_scope,
                                benchmark_mode
                            )
                        )

                        st.success(
                            "Results saved to evaluation_results.json "
                            "and evaluation_results.csv."
                        )

                    except Exception as e:

                        st.session_state.benchmark_report = {
                            "generated_at": datetime.now().isoformat(
                                timespec="seconds"
                            ),
                            "source_scope": benchmark_scope,
                            "benchmark_mode": benchmark_mode,
                            "run_status": benchmark_summary["run_status"],
                            "successful_cases": benchmark_summary[
                                "successful_cases"
                            ],
                            "failed_cases": benchmark_summary["failed_cases"],
                            "quota_exhausted": benchmark_summary[
                                "quota_exhausted"
                            ],
                            "summary": benchmark_summary,
                            "cases": benchmark_results
                        }

                        st.warning(
                            f"Results could not be saved to disk: {e}"
                        )

            benchmark_report = (
                st.session_state.get("benchmark_report")
                or load_saved_benchmark_report()
            )

            if benchmark_report:

                report_summary = benchmark_report["summary"]

                st.markdown("#### 📊 Benchmark Results")

                report_successful_cases = report_summary.get(
                    "successful_cases",
                    report_summary.get("total_cases", 0)
                    - report_summary.get("failed_cases", 0)
                )

                report_quota_exhausted = report_summary.get(
                    "quota_exhausted",
                    False
                )

                st.caption(
                    f"Last run: {benchmark_report.get('generated_at', '—')} · "
                    f"{benchmark_report.get('benchmark_mode', 'Benchmark')} · "
                    f"Source: {benchmark_report.get('source_scope', '—')} · "
                    f"{report_summary.get('total_cases', 0)} cases: "
                    f"{report_successful_cases} successful, "
                    f"{report_summary.get('failed_cases', 0)} failed, "
                    f"{report_summary.get('skipped_cases', 0)} skipped"
                )

                show_benchmark_scores = report_successful_cases > 0

                if not show_benchmark_scores:

                    if report_quota_exhausted:
                        st.error(
                            "Benchmark stopped — Gemini API quota "
                            "reached. No valid scores were produced. "
                            + QUOTA_ERROR_MESSAGE
                        )
                    else:
                        st.error(
                            "Benchmark stopped — no case completed "
                            "successfully. No valid scores were produced."
                        )

                elif report_quota_exhausted:

                    st.warning(
                        "Partial benchmark: Gemini API quota was reached. "
                        f"Scores below are based on only "
                        f"{report_successful_cases} successful case(s)."
                    )

                if show_benchmark_scores:

                    avg1, avg2, avg3, avg4 = st.columns(4)

                    with avg1:
                        st.metric(
                            "Average Faithfulness",
                            format_metric(
                                report_summary.get("average_faithfulness")
                            )
                        )

                    with avg2:
                        st.metric(
                            "Average Answer Relevancy",
                            format_metric(
                                report_summary.get("average_answer_relevancy")
                            )
                        )

                    with avg3:
                        st.metric(
                            "Average Context Precision",
                            format_metric(
                                report_summary.get("average_context_precision")
                            )
                        )

                    with avg4:
                        st.metric(
                            "Average Context Recall",
                            format_metric(
                                report_summary.get("average_context_recall")
                            )
                        )

                    acc1, acc2 = st.columns(2)

                    with acc1:
                        st.metric(
                            "Refusal Accuracy",
                            format_metric(
                                report_summary.get("refusal_accuracy"),
                                as_percent=True
                            ),
                            help=(
                                "Share of cases where the tutor refused "
                                "exactly when it should: off-material "
                                "questions refused, in-material "
                                "questions answered."
                            )
                        )
                        st.caption(
                            f"{report_summary.get('refusal_correct', 0)} / "
                            f"{report_summary.get('refusal_total', 0)} correct"
                        )

                    with acc2:
                        st.metric(
                            "Citation / Expected Source Match",
                            format_metric(
                                report_summary.get("source_match_rate"),
                                as_percent=True
                            ),
                            help=(
                                "Expected citation found among the "
                                "retrieved sources or in the answer."
                            )
                        )
                        st.caption(
                            f"{report_summary.get('source_match_correct', 0)} / "
                            f"{report_summary.get('source_match_total', 0)} matched"
                        )

                def yes_no_mark(value):

                    if value is None:
                        return "—"

                    return "✅" if value else "❌"

                result_rows = []

                for case_result in benchmark_report["cases"]:

                    if case_result.get("should_refuse"):

                        refusal_text = {
                            True: "Correct refusal ✅",
                            False: "Incorrect ❌"
                        }.get(
                            case_result.get("refusal_correct"),
                            "—"
                        )

                    else:

                        refusal_text = {
                            True: "Answered ✅",
                            False: "False refusal ❌"
                        }.get(
                            case_result.get("refusal_correct"),
                            "—"
                        )

                    result_rows.append({
                        "ID": case_result.get("id", ""),
                        "Question": case_result.get("question", ""),
                        "Status": case_result.get("status", ""),
                        "Faithfulness": case_result.get("faithfulness"),
                        "Answer Relevancy": case_result.get(
                            "answer_relevancy"
                        ),
                        "Context Precision": case_result.get(
                            "context_precision"
                        ),
                        "Context Recall": case_result.get(
                            "context_recall"
                        ),
                        "Refusal": refusal_text,
                        "Expected Source": (
                            case_result.get("expected_source") or "—"
                        ),
                        "Source Match": yes_no_mark(
                            case_result.get("source_match")
                        ),
                        "Error": case_result.get("error", "")
                    })

                st.dataframe(
                    result_rows,
                    hide_index=True
                )

                with st.expander("🔎 Benchmark answers and retrieved sources"):

                    for case_result in benchmark_report["cases"]:

                        st.markdown(
                            f"**{case_result.get('id', '')} — "
                            f"{case_result.get('question', '')}**"
                        )

                        st.write(
                            case_result.get("answer") or "—"
                        )

                        st.caption(
                            "Retrieved: "
                            + (
                                "; ".join(
                                    case_result.get(
                                        "retrieved_citations",
                                        []
                                    )
                                )
                                or "—"
                            )
                        )

                        st.divider()

            else:

                st.info(
                    "No benchmark results yet. Choose Quick Benchmark "
                    "and click Run Benchmark — results appear here and "
                    "are saved to evaluation_results.json / .csv."
                )

    # -------------------------------------------
    # QUESTION REPETITION RATE
    # -------------------------------------------

    st.divider()

    st.subheader("🔁 Question Repetition Rate")

    st.caption(
        "Each accepted assessment question is compared with "
        "all earlier ones using the same normalization and "
        "similarity rule as duplicate prevention "
        "(SequenceMatcher ratio ≥ 0.82)."
    )

    question_repetition = compute_question_repetition_rate(
        st.session_state.get("quiz_history", [])
    )

    if question_repetition["total_questions"] < 2:

        st.info("Not enough assessment history yet.")

    else:

        st.metric(
            "Question Repetition Rate",
            format_repetition_rate(question_repetition)
        )

    # -------------------------------------------
    # PERSONALIZATION BENCHMARK (SIMULATED)
    # -------------------------------------------

    st.divider()

    with st.container(key="card_synthetic"):

        section_header(
            "🧑‍🎓 Personalization Benchmark — Synthetic",
            "A deterministic simulation of the learner model, not "
            "real-student evidence.",
            eyebrow="Synthetic · simulated learners"
        )

        st.badge(
            "SYNTHETIC — simulated profiles, not real students",
            icon=":material/science:",
            color="violet"
        )

        st.warning(
            "Synthetic simulation — these are NOT real student "
            "results. Three simulated learner profiles are replayed "
            "through the same EMA mastery update used by LearnLens "
            f"(quiz evidence 1.0 / 0.0, α = {QUIZ_MASTERY_ALPHA:.2f}) "
            "across two sessions. Deterministic, no API calls."
        )

        st.dataframe(
            run_personalization_benchmark(),
            hide_index=True
        )

        if question_repetition["total_questions"] < 2:

            repetition_summary_text = (
                "Not enough assessment history yet."
            )

        else:

            repetition_summary_text = format_repetition_rate(
                question_repetition
            )

        st.markdown(
            "**Question repetition rate (real assessment "
            f"history):** {repetition_summary_text}"
        )


# ===================================================
# SIDEBAR - SESSION STATUS
# (rendered last so it reflects this run's uploads;
# no API calls)
# ===================================================

with st.sidebar:

    st.markdown("### 🎓 Learning Session")

    status_items, _ = build_unified_knowledge_base()

    status_counts = {
        "pdf": 0,
        "slide": 0,
        "video": 0
    }

    for status_item in status_items:
        status_counts[status_item["source_type"]] = (
            status_counts.get(status_item["source_type"], 0) + 1
        )

    # -------------------------------------------
    # SESSION
    # -------------------------------------------

    eyebrow_label("Session")

    st.metric(
        "Knowledge units",
        len(status_items)
    )

    side_pdf, side_slides, side_video = st.columns(3)

    with side_pdf:
        st.caption("PDF")
        st.markdown(f"**{status_counts['pdf']}**")

    with side_slides:
        st.caption("Slides")
        st.markdown(f"**{status_counts['slide']}**")

    with side_video:
        st.caption("Video")
        st.markdown(f"**{status_counts['video']}**")

    if not status_items:
        st.caption(
            "Upload learning materials to start."
        )

    st.divider()

    # -------------------------------------------
    # LEARNER
    # -------------------------------------------

    eyebrow_label("Learner")

    status_profile = st.session_state.get(
        "learner_profile",
        {}
    )

    side_level, side_difficulty = st.columns(2)

    with side_level:

        st.metric(
            "Level",
            status_profile.get("level", "Intermediate")
        )

    with side_difficulty:

        st.metric(
            "Next difficulty",
            st.session_state.get("quiz_difficulty", "Medium")
        )

    if status_profile.get("diagnostic_completed"):

        st.badge(
            "Diagnostic "
            f"{status_profile.get('diagnostic_score', 0)}/"
            f"{status_profile.get('diagnostic_total', 3)}",
            icon=":material/check_circle:",
            color="green"
        )

    else:

        st.badge(
            "Diagnostic not completed",
            icon=":material/pending:",
            color="gray"
        )

    st.divider()

    # -------------------------------------------
    # KNOWLEDGE
    # -------------------------------------------

    eyebrow_label("Knowledge")

    status_tagged = sum(
        1 for status_item in status_items
        if status_item.get("tagged")
    )

    st.caption(
        f"Tagged units: {status_tagged} / {len(status_items)}"
    )

    st.progress(
        status_tagged / len(status_items)
        if status_items
        else 0.0
    )

    st.caption(
        "Topics tracked in learner model: "
        f"{len(st.session_state.get('learner_mastery', {}))}"
    )
